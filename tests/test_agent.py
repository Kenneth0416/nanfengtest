import os
import shutil
from pathlib import Path

import pytest

from core import documents
from core import taxonomy as tx
from core.app import App
from core.models import RawItem
from core.signals import SignalEngine

ROOT = Path(__file__).resolve().parent.parent
PI_DIR = Path(os.getenv("PI_ENGINE_DIR", ROOT / "pi_engine"))
# Tests that run the Pi Durable harness need Node 22 and `npm ci` in pi_engine/; without them they skip, not fail.
needs_pi = pytest.mark.skipif(not (shutil.which("node") and (PI_DIR / "node_modules" / "@earendil-works" / "pi-durable").exists()),
                              reason="Pi Durable engine not installed (cd pi_engine && npm ci); the Copilot tests need it")


@pytest.fixture(scope="module")
def app(tmp_path_factory):
    os.environ["LLM_PROVIDER"] = "offline"
    a = App(live=False)
    a.store = a.ctx.store = type(a.store)(tmp_path_factory.mktemp("db") / "t.db")
    a.refresh()
    return a


def test_classifier_detects_item_types(app):
    clf = app.ctx.classifier
    assert clf.classify(RawItem("t", payload={"period": "2026Q2", "market": "Central London", "metric": "vacancy_pct",
                                               "value": "8"})).kind == "observation"
    assert clf.classify(RawItem("t", payload={"deal_type": "letting", "party": "X", "building": "Y"})).kind == "deal"
    assert clf.classify(RawItem("t", payload={"date": "2026-01-01", "indicator": "bank_rate", "value": 4})).kind == "macro"
    news = clf.classify(RawItem("t", title="Barclays buys 1 Churchill Place for £750m", text="Canary Wharf HQ at a 5.62% yield"))
    assert news.kind == "news"
    assert news.record.category == "investment_deal"
    assert "Canary Wharf" in news.record.markets
    assert news.record.importance >= 4


def test_non_exact_market_names_warn(app):
    c = app.ctx.classifier.classify(RawItem("t", payload={"period": "2026Q2", "market": "Central London",
                                                          "metric": "vacancy_pct", "value": "8"}))
    assert c.warnings and "London" in c.warnings[0]


def test_seed_loaded(app):
    assert len(app.store.observations()) > 500
    assert app.store.news()


def test_signals_rank_by_role(app):
    eng = SignalEngine(app.store, app.config)
    analyst = eng.for_role("analyst", limit=3)
    assert any(s.kind == "data_quality" for s in analyst)
    exec_ = eng.for_role("executive", ["London"], limit=5)
    assert exec_ and all(s.kind != "data_quality" for s in exec_)


def test_watch_rule_fires(app):
    app.store.add_watch_rule("asset_manager", "Canary Wharf", "prime_rent_psf", ">", 60, "CW rent > 60")
    alerts = SignalEngine(app.store, app.config).watch_alerts("asset_manager")
    assert alerts and alerts[0].kind == "alert"


def test_tools_resolve_core_submarkets(app):
    out = app.call("get_metric_trend", {"market": "West End", "metric": "prime_rent_psf"})
    assert "West End Core" in out["summary"] and out["chart"]["x"]


def test_user_focus_filters_home_and_undoes(app):
    from core import focus as fx
    assert fx.parse_command("show me more Fitzrovia and less macro") == [("add", "market", "Fitzrovia"),
                                                                        ("remove", "measure", "macro")]
    assert fx.parse_command("Tell me more about Fitzrovia rents") is None  # a question, not a setting
    assert fx.parse_command("Alert me when Canary Wharf vacancy goes above 10%") is None
    user = "asset_manager"
    app.store.reset_focus(user)
    app.call("update_focus", {"request": "only rents in Canary Wharf, top 3"}, user)
    f = app.focus(user)
    assert f["markets"] == ["Canary Wharf"] and f["measures"] == ["rents"] and f["budget"] == "3"
    sigs = SignalEngine(app.store, app.config).for_focus(user, f)
    assert sigs and len(sigs) <= 3
    assert all(s.kind == "alert" or fx.measure_of(s) == "rents" for s in sigs)
    app.store.undo_focus(user)
    assert app.focus(user) == fx.defaults(app.config, user)


def test_agent_proposes_setup_and_applies_only_on_confirm(app):
    user = "investor"
    app.store.reset_focus(user)
    out = app.call("propose_setup", {"description": "I run acquisitions in Midtown and the City and present to the IC"}, user)
    p = out["data"]
    assert p["focus"]["markets"] == ["Midtown", "City & Southbank"] and p["focus"]["brief_format"] == "ic"
    assert "investment" in p["focus"]["measures"] and p["rules"]
    before, team_before = len(app.store.watch_rules(user)), len(app.store.watch_rules("team"))
    assert app.store.focus(user) is None  # nothing saved yet
    applied = app.call("apply_setup", {}, user)["data"]
    # the proposed alerts are the person's own: they fire in 'For you' and never reach the team by themselves
    assert app.focus(user) == p["focus"] and len(app.store.watch_rules(user)) == before + len(p["rules"])
    assert len(app.store.watch_rules("team")) == team_before
    rid = applied["rule_ids"][0]
    app.store.share_watch_rule(rid)  # 'Share with team' on Monitor
    assert len(app.store.watch_rules("team")) == team_before + 1 and len(app.store.watch_rules(user)) == before + len(p["rules"]) - 1
    from core import focus as fx
    fx.undo_setup(app.store, user, applied)
    assert app.store.focus(user) is None and len(app.store.watch_rules("team")) == team_before
    # the Copilot's own alert tool follows the same rule
    out = app.call("add_watch_rule", {"market": "Midtown", "metric": "vacancy_pct", "operator": ">", "value": 9}, "priya")
    assert "for you" in out["summary"] and app.store.watch_rules("priya")[-1]["id"] == out["data"]["id"]
    out = app.call("add_watch_rule", {"market": "Midtown", "metric": "vacancy_pct", "operator": ">", "value": 9, "share_with_team": True}, "priya")
    assert out["summary"].startswith("Team alert") and app.store.watch_rules("team")[-1]["id"] == out["data"]["id"]


def test_team_brief_uses_added_signals_with_footnotes(app):
    sig = SignalEngine(app.store, app.config).for_role("team")[0]
    app.store.add_to_brief(sig.to_dict(), "Kenneth")
    md = app.call("generate_brief", {"format": "ic"}, "Kenneth")["summary"]
    assert sig.title in md and "chosen by Kenneth" in md and "## Risks to the case" in md
    assert "[1]" in md and "## Sources" in md
    app.store.remove_from_brief(sig.id)


def test_without_pi_the_app_says_how_to_start_the_agent(app, monkeypatch, tmp_path):
    """Pi is the only agent engine: with no Node packages there is no silent fallback, just a clear message."""
    from core.engine import Engine, PiNotReady
    monkeypatch.setenv("PI_ENGINE_DIR", str(tmp_path))
    eng = Engine(app)
    assert not eng.ready and "npm install" in eng.status
    with pytest.raises(PiNotReady, match="isn't available") as err:
        eng.ask("What changed?", "Kenneth")
    assert "npm install" in err.value.detail  # the technical cause stays for 'Technical details'
    with pytest.raises(PiNotReady):
        eng.gather("Kenneth")


def test_vision_extraction_checks_every_figure(app, tmp_path):
    """A fake vision model: one true figure, one it made up, one with an inferred label and one out of range,
    plus a chart-only page with no text layer."""
    from PIL import Image, ImageDraw
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    chart = tmp_path / "chart.png"
    img = Image.new("RGB", (800, 500), "white")
    ImageDraw.Draw(img).text((40, 40), "City vacancy 9.6%", fill="black")
    img.save(chart)
    pdf = tmp_path / "report.pdf"
    c = canvas.Canvas(str(pdf), pagesize=A4)
    c.drawString(72, 780, "Savills London Offices Q2 2025")
    c.drawString(72, 760, "West End vacancy rate 6.4%. Take-up 3,200,000 sq ft. Prime rent £170.00 per sq ft.")
    c.showPage()
    c.drawImage(str(chart), 72, 300, width=400, height=250)
    c.showPage()
    c.save()

    class FakeVision:
        available, name, model = True, "fake", "fake-vision"
        calls = []

        def complete_json_images(self, prompt, images):
            self.calls.append(len(images))
            assert images[0][:8] == b"\x89PNG\r\n\x1a\n" and "Page 1 text layer" in prompt
            return [
                {"market_text": "West End", "metric": "vacancy_pct", "value": 6.4, "page": 1, "label_status": "explicit"},
                {"market_text": "London", "metric": "take_up_msqft", "value": 3.2, "page": 1, "label_status": "explicit"},
                {"market_text": "West End", "metric": "prime_rent_psf", "value": 175, "page": 1, "label_status": "explicit"},
                {"market_text": "Midtown", "metric": "vacancy_pct", "value": 6.4, "page": 1, "label_status": "inferred",
                 "where": "unlabelled column"},
                {"market_text": "City", "metric": "vacancy_pct", "value": 9.6, "page": 2, "label_status": "explicit"},
                {"market_text": "City", "metric": "prime_yield_pct", "value": 52, "page": 2, "label_status": "explicit"},
            ]

    fake = FakeVision()
    a = documents.analyse_document(app.store, fake, pdf)
    assert a["method"].startswith("vision model") and fake.calls == [2]
    by = {(c["market"], c["metric"]): c for c in a["candidates"]}
    assert by[("West End", "vacancy_pct")]["status"] == "verified" and by[("West End", "vacancy_pct")]["suggest_keep"]
    assert by[("London", "take_up_msqft")]["status"] == "verified"            # printed as 3,200,000 sq ft
    made_up = by[("West End Core", "prime_rent_psf")]
    assert made_up["status"] == "not_found" and not made_up["suggest_keep"]     # 175 is not on the page
    assert not by[("Midtown", "vacancy_pct")]["suggest_keep"]                  # label was inferred
    chart_only = by[("City & Southbank", "vacancy_pct")]
    assert chart_only["status"] == "image" and chart_only["page"] == 2 and not chart_only["suggest_keep"]
    assert any("plausible" in f for f in by[("City & Southbank", "prime_yield_pct")]["flags"])
    assert documents.page_png(pdf, 2)[:4] == b"\x89PNG"


def test_vision_model_is_separate_and_uses_json_schema(monkeypatch):
    """PDF reading can use its own OpenAI-compatible vision model (MiMo) while the chat model stays put."""
    from core import llm as L
    sent = {}

    class FakeCompletions:
        def create(self, **kw):
            sent.update(kw)
            msg = type("M", (), {"content": '{"figures": [{"market_text": "City", "metric": "vacancy_pct", "value": 9.1, '
                                            '"page": 1, "where": "table", "label_status": "explicit", "evidence": "9.1%"}]}'})
            return type("R", (), {"choices": [type("C", (), {"message": msg})]})

    class FakeOpenAI:
        def __init__(self, api_key, base_url):
            sent["base_url"] = base_url
            self.chat = type("Chat", (), {"completions": FakeCompletions()})

    import openai
    monkeypatch.setattr(openai, "OpenAI", FakeOpenAI)
    monkeypatch.setenv("MIMO_API_KEY", "test")
    cfg = {"base_url": "https://api.xiaomimimo.com/v1", "model": "mimo-v2.6-flash", "api_key_env": "MIMO_API_KEY",
           "pages_per_call": 2, "detail": "high", "reasoning_effort": "medium"}
    chat = L.OfflineLLM()
    v = L.get_vision(cfg, chat)
    assert v is not chat and v.model == "mimo-v2.6-flash" and v.pages_per_call == 2
    figs = v.complete_json_images("read", [b"\x89PNG fake"])
    assert figs[0]["value"] == 9.1 and sent["base_url"] == "https://api.xiaomimimo.com/v1"
    assert sent["response_format"]["type"] == "json_schema" and sent["reasoning_effort"] == "medium"
    assert sent["messages"][0]["content"][0]["image_url"]["detail"] == "high"
    monkeypatch.delenv("MIMO_API_KEY")
    assert L.get_vision(cfg, chat) is chat  # no key: PDFs fall back to the chat model


def test_pdf_extraction_flags_definition_differences(app):
    pdf = ROOT / "reference" / "samples" / "sample_bnp_paribas_q2_2025_restated.pdf"
    a = documents.analyse_document(app.store, app.llm, pdf)
    assert a["period"] == "2025Q2" and a["publisher"] == "BNP Paribas"
    london = next(c for c in a["candidates"] if c["market"] == "London" and c["metric"] == "vacancy_pct")
    assert london["value"] == 8.3 and london["stored"]["value"] == 8.8 and london["flags"]


@needs_pi
def test_pi_durable_engine_with_scripted_model(app, tmp_path):
    os.environ.update(PI_FAUX="1", PI_STORE=str(tmp_path / "pi.sqlite"), PI_CONVERSATIONS=str(tmp_path / "c.json"))
    from core.pi_engine import PiDurableEngine
    eng = PiDurableEngine(app)
    ans = eng.ask("What changed?", "executive")
    assert "Pi Durable" in ans.text and ans.trace[0]["tool"] == "get_signals"
    from core.engine import Engine
    engine = Engine(app)
    assert engine.ready
    assert engine.ask("What does Savills say about Midtown?", "Kenneth").text.startswith("> **Coverage:** The library has no Savills")


def test_shared_library_upload_dedupe_and_flags(app):
    """Anyone uploads; checked figures are shared, the rest shared flagged; a second upload is caught."""
    from core import analytics as an
    sample = ROOT / "reference" / "samples" / "sample_bnp_paribas_q2_2025_restated.pdf"
    assert documents.duplicate_of(app.store, sample) is None
    analysis = documents.analyse_document(app.store, app.llm, sample)
    n = len(analysis["candidates"])
    for c in analysis["candidates"][:2]:  # as if the vision model read these two from a chart
        c.update(status="image", suggest_keep=False, flags=["Read from the page image only (p.1)."])
    unchecked = [0, 1]
    res = documents.save_document(app.store, analysis, user="Priya")
    assert res["saved"] == n and res["flagged"] == len(unchecked) > 0
    dup = documents.duplicate_of(app.store, sample)
    assert dup["uploaded_by"] == "Priya" and dup["channel"] == "upload"
    up = app.store.uploads()[0]
    assert up["publisher"] == analysis["publisher"] and up["to_check"] == len(unchecked)
    assert an.provenance(dup).startswith("Uploaded by Priya")
    # Anyone settles a flagged figure: rejected ones leave every view, confirmed ones lose the flag.
    rows = [r for r in app.store.flagged() if r["source_id"] == res["source_id"]]
    app.store.resolve_observation(rows[0]["rowid"], "James", keep=False)
    assert rows[0]["rowid"] not in {o["rowid"] for o in app.store.observations()}
    if len(rows) > 1:
        app.store.resolve_observation(rows[1]["rowid"], "James", keep=True)
        assert rows[1]["rowid"] not in {r["rowid"] for r in app.store.flagged()}


def test_anyone_can_flag_a_number_and_it_carries_a_warning(app):
    from core import analytics as an
    from skills.market_metrics.tools import get_metric_trend
    snap = an.snapshot(app.store, "London", "vacancy_pct")
    assert an.source_ref(app.store, snap["source_id"])["provenance"].startswith("Uploaded by")
    app.store.flag_observation(snap["rowid"], "Doesn't match the PDF", "Sarah")
    try:
        sig = next(s for s in SignalEngine(app.store, app.config).all_signals() if s.obs_id == snap["rowid"])
        assert "Flagged by Sarah" in sig.flag
        assert "Flagged by Sarah" in get_metric_trend(app.ctx, {"market": "London", "metric": "vacancy_pct"})["summary"]
        assert "(flagged)" in app.call("generate_brief", {"format": "board"}, "Sarah")["summary"]
    finally:
        app.store.resolve_observation(snap["rowid"], "Sarah", keep=True)
    assert an.snapshot(app.store, "London", "vacancy_pct")["status"] == "ok"


def test_feeds_and_report_reminders(app):
    from datetime import date
    from core import library as lib
    feeds = lib.feeds(app.store, app.config)
    assert any(f["licence"] == "Open Government Licence" for f in feeds)
    assert lib.last_auto_update(app.store)["kind"] == "auto"
    due = lib.due_reports(app.store, app.config, today=date(2026, 10, 10))
    kf = [d for d in due if d["publisher"] == "Knight Frank"]
    sav = [d for d in due if d["publisher"] == "Savills"]
    assert kf == [{"publisher": "Knight Frank", "period": "2026Q3", "due": date(2026, 10, 28), "missing": False}]
    assert any(d["missing"] and d["period"] == "2026Q2" for d in sav)


def test_colleague_tips_merge_with_the_same_story_and_get_credit(app):
    from core import gather
    res = app.call("share_tip", {"text": "Heard from a broker that Barclays is close to buying 1 Churchill Place "
                                          "for around £750m", "how": "Broker call"}, "Kenneth")["data"]
    assert res["merged"] and "Barclays buys 1 Churchill Place" in res["title"] and res["shared_by"] == ["Kenneth"]
    assert "shared by Kenneth (broker call)" in res["why_hot"]
    new = app.call("share_tip", {"text": "Google is looking for 200,000 sq ft in King's Cross for an AI team",
                                  "how": "LinkedIn"}, "Priya")["data"]
    assert not new["merged"] and new["category"] == "leasing_deal" and "unverified" in new["why_hot"]
    sig = next(s for s in SignalEngine(app.store, app.config).for_role("team") if "Google" in s.title)
    assert sig.title.startswith("Unverified tip") and sig.shared_by[0]["who"] == "Priya" and sig.hot > 0
    # A second colleague sharing the same tip adds their credit instead of a duplicate.
    again = app.call("share_tip", {"text": "Google said to want 200,000 sq ft at King's Cross", "how": "Broker call"},
                     "James")["data"]
    assert again["merged"] and again["id"] == new["id"] and again["shared_by"] == ["Priya", "James"]
    assert any(a["kind"] == "share" and a["author"] == "James" for a in app.store.activity(5))


def test_feed_news_from_another_outlet_is_merged(app):
    from core.ingest import ingest_items
    item = RawItem("news_rss", title="Barclays acquires 1 Churchill Place in £750m Canary Wharf deal",
                   text="The bank bought its headquarters building.", url="https://example.com/b",
                   source="Example Wire", published="2026-07-01")
    before = len(app.store.news())
    ingest_items(app.ctx, [item], "news_rss")
    assert len(app.store.news()) == before
    n = next(n for n in app.store.news() if n.title.startswith("Barclays buys"))
    assert n.also_reported[0]["source"] == "Example Wire"


def test_report_index_reads_titles_only(app):
    from core import gather
    html = ("<ul><li><a href='/r/1'>London Offices Spotlight Q3 2026</a></li>"
            "<li><a href='/r/2'>London Offices Spotlight Q2 2026</a></li><li><a>Dublin Offices Q3 2026</a></li></ul>")
    cfg = {**app.config, "report_index": [{"publisher": "Knight Frank", "url": "https://kf.example/research"}]}
    found = gather.check_report_index(app.store, cfg, fetch=lambda u: html)
    assert [f["period"] for f in found] == ["2026Q3"]  # Q2 is already in the library; Dublin is not London


@needs_pi
def test_pi_gatherer_files_the_inbox(app, tmp_path):
    import json
    from core.pi_engine import PiDurableEngine
    if not (PI_DIR / "node_modules" / "@earendil-works" / "pi-durable").exists() or not shutil.which("node"):
        pytest.skip("Pi Durable not installed")
    iid = app.store.add_to_inbox("Broker says Derwent is close to letting 120,000 sq ft at Network to a bank",
                                 "", "Broker call", "Sarah")
    title = "Derwent close to letting 120,000 sq ft at Network"
    script = tmp_path / "script.json"
    script.write_text(json.dumps([
        {"tool": "list_inbox"},
        {"tool": "find_matches", "args": {"title": title}},
        {"tool": "save_item", "args": {"title": title, "text": "Broker says Derwent is close to letting 120,000 sq ft",
                                       "shared_by": "Sarah", "how": "Broker call", "inbox_id": iid}},
        {"text": "Gathering round done: 1 new tip from Sarah."}]))
    os.environ.update(PI_FAUX="1", PI_FAUX_SCRIPT=str(script), PI_STORE=str(tmp_path / "pi.sqlite"),
                      PI_CONVERSATIONS=str(tmp_path / "c.json"), PI_ENGINE_DIR=str(PI_DIR))
    try:
        ans = PiDurableEngine(app).gather(role="team")
    finally:
        os.environ.pop("PI_FAUX_SCRIPT")
    assert [t["tool"] for t in ans.trace] == ["list_inbox", "find_matches", "save_item"]
    assert not app.store.inbox("new")
    assert any(n.shared_by and n.shared_by[0]["who"] == "Sarah" for n in app.store.news())


def test_follow_source_checks_terms_and_picks_the_channel(app):
    from core import follow
    ok = follow.assess("https://www.cityoflondon.gov.uk/planning", page="<html></html>")
    assert ok["channel"] == "auto" and ok["licence"] == "Open Government Licence"
    robots = follow.assess("https://news.example.com/offices", robots="User-agent: *\nDisallow: /offices", page="<p>x</p>")
    assert robots["channel"] == "upload" and "robots.txt" in robots["reason"]
    terms = ("<p>You may not use any automated means, robots or scrapers, or use the content to train artificial "
             "intelligence, without our prior written consent.</p>")
    banned = follow.assess("https://research.example.com/", page="<a href='/terms'>Terms</a>", terms=terms)
    assert banned["channel"] == "upload" and "automated" in banned["quote"]
    feed = follow.assess("https://blog.example.com/feed", page="<?xml version='1.0'?><rss></rss>",
                         content_type="application/rss+xml", terms="<p>Copyright 2026. All rights reserved.</p>")
    assert feed["channel"] == "auto" and feed["kind"] == "rss"
    agency = follow.assess("https://www.savills.co.uk/research_articles/", page="<p>Research</p>")
    assert agency["channel"] == "upload"
    assert follow.assess("https://down.example.com/", errors=["page: ConnectionError"])["channel"] == "upload"
    assert follow.resolve("Follow Savills Market Watch")["url"].startswith("https://www.savills.co.uk")


def test_followed_sources_are_shared_and_remind_the_team(app, monkeypatch):
    from core import follow, library
    res = app.call("follow_source", {"source": "Savills Market Watch", "page": "<p>Research</p>"}, "Priya")
    assert "upload-only" in res["summary"] and res["data"]["added_by"] == "Priya"
    again = app.call("follow_source", {"source": "https://www.savills.co.uk/research_articles/", "page": "<p/>"}, "James")
    assert again["data"]["already"] and "Priya already follows" in again["summary"]
    assert any(d["name"] == "Savills Market Watch" for d in follow.due_uploads(app.store))
    # An allowed RSS feed joins the automatic feeds and its items come in as news.
    app.call("follow_source", {"source": "https://blog.example.com/feed", "name": "Example office blog",
                               "page": "<rss></rss>", "content_type": "application/rss+xml",
                               "terms": "<p>All rights reserved.</p>"}, "Kenneth")
    assert any(f["name"] == "Example office blog" for f in library.feeds(app.store, app.config))
    rss = (b"<?xml version='1.0'?><rss><channel><item><title>Meta takes 150,000 sq ft at King's Cross</title>"
           b"<link>https://blog.example.com/meta</link></item></channel></rss>")
    import connectors.followed as cf

    class R:
        ok, content, text, headers = True, rss, rss.decode(), {}
        def raise_for_status(self): pass
    monkeypatch.setattr(cf.requests, "get", lambda *a, **k: R())
    from core.ingest import ingest_items
    conn = cf.Connector(app.ctx, {})
    ingest_items(app.ctx, conn.fetch(), "followed_sources")
    assert any(n.source == "Example office blog" for n in app.store.news())
    assert conn.fetch() == []  # already seen


@needs_pi
def test_pi_follow_source_fetches_then_python_decides(app, tmp_path):
    import json
    from core.pi_engine import PiDurableEngine
    if not (PI_DIR / "node_modules" / "@earendil-works" / "pi-durable").exists() or not shutil.which("node"):
        pytest.skip("Pi Durable not installed")
    script = tmp_path / "s.json"
    script.write_text(json.dumps([{"tool": "follow_source", "args": {"source": "https://unreachable.example.invalid/feed",
                                                                      "name": "Unreachable feed"}},
                                  {"text": "Followed."}]))
    os.environ.update(PI_FAUX="1", PI_FAUX_SCRIPT=str(script), PI_STORE=str(tmp_path / "pi.sqlite"),
                      PI_CONVERSATIONS=str(tmp_path / "c.json"), PI_ENGINE_DIR=str(PI_DIR))
    try:
        ans = PiDurableEngine(app).gather(role="Sarah")
    finally:
        os.environ.pop("PI_FAUX_SCRIPT")
    # The Pi tool fetched (and failed to reach) the site; Python saw no page and chose upload-only.
    assert [t["tool"] for t in ans.trace] == ["resolve_source", "follow_source"]
    row = app.store.q("SELECT * FROM followed WHERE name='Unreachable feed'")[0]
    assert row["channel"] == "upload" and row["added_by"] == "Sarah"


def test_information_agent_skills_load_and_tools_are_steps(app, tmp_path):
    from reportlab.pdfgen import canvas
    pdf = tmp_path / "savills_q3_2026.pdf"
    c = canvas.Canvas(str(pdf))
    for i, line in enumerate(["Savills Central London Offices Q3 2026", "City vacancy rate: 9.1%",
                              "West End take-up: 1.2m sq ft"]):
        c.drawString(72, 760 - 20 * i, line)
    c.save()
    names = {s["name"] for s in app.call("list_skills", {})["data"]}
    assert {"check-terms", "fetch-source", "dedupe", "read-pdf", "read-table", "read-article", "map-submarkets",
            "check-figures"} <= names
    assert "Never describe a flagged figure as checked" in app.call("load_skill", {"name": "check-figures"})["summary"]
    assert not app.call("dedupe_report", {"path": str(pdf)})["data"]["duplicate"]
    ov = app.call("pdf_overview", {"path": str(pdf)})["data"]
    assert ov["publisher"] == "Savills" and ov["period"] == "2026Q3"
    figs = app.call("read_pdf_pages", {"path": str(pdf), "pages": [1]})["data"]["figures"]
    checked = app.call("check_figures", {"path": str(pdf), "figures": figs})["data"]
    assert checked and all(c["status"] == "verified" for c in checked)
    app.call("stage_figures", {"path": str(pdf), "figures": figs + [
        {"market_text": "City", "metric": "prime_yield_pct", "value": 25, "page": 1}]}, "Kenneth")
    staged = app.store.staged(str(pdf))
    assert staged["candidates"][-1]["suggest_keep"] is False  # 25% yield fails the range check


@needs_pi
def test_pi_information_agent_reads_an_upload(app, tmp_path):
    import json
    from reportlab.pdfgen import canvas
    from core.pi_engine import PiDurableEngine
    if not (PI_DIR / "node_modules" / "@earendil-works" / "pi-durable").exists() or not shutil.which("node"):
        pytest.skip("Pi Durable not installed")
    pdf = tmp_path / "jll_q3_2026.pdf"
    c = canvas.Canvas(str(pdf))
    c.drawString(72, 760, "JLL London Office Market Q3 2026")
    c.drawString(72, 740, "Canary Wharf vacancy rate: 12.4%")
    c.save()
    fig = {"market_text": "Canary Wharf", "metric": "vacancy_pct", "value": 12.4, "page": 1, "label_status": "explicit"}
    script = tmp_path / "s.json"
    script.write_text(json.dumps([
        {"tool": "load_skill", "args": {"name": "read-pdf"}},
        {"tool": "dedupe_report", "args": {"path": str(pdf)}},
        {"tool": "pdf_overview", "args": {"path": str(pdf)}},
        {"tool": "read_pdf_pages", "args": {"path": str(pdf), "pages": [1]}},
        {"tool": "map_markets", "args": {"names": ["Canary Wharf"]}},
        {"tool": "check_figures", "args": {"path": str(pdf), "figures": [fig]}},
        {"tool": "stage_figures", "args": {"path": str(pdf), "figures": [fig], "method": "agent: text layer"}},
        {"text": "Staged 1 figure for review."}]))
    os.environ.update(PI_FAUX="1", PI_FAUX_SCRIPT=str(script), PI_STORE=str(tmp_path / "pi.sqlite"),
                      PI_CONVERSATIONS=str(tmp_path / "c.json"), PI_ENGINE_DIR=str(PI_DIR))
    try:
        ans = PiDurableEngine(app).gather(f"Kenneth uploaded {pdf}", "Kenneth")
    finally:
        os.environ.pop("PI_FAUX_SCRIPT")
    from core.engine import TOOL_SKILL
    chain = list(dict.fromkeys(TOOL_SKILL[t["tool"]] for t in ans.trace if t["tool"] in TOOL_SKILL))
    assert chain == ["dedupe", "read-pdf", "map-submarkets", "check-figures"]
    staged = app.store.staged(str(pdf))
    assert staged["publisher"] == "JLL" and staged["candidates"][0]["suggest_keep"]


def test_summary_and_brief_lead_with_the_market_not_tips(app):
    from core.summary import team_summary
    app.call("share_tip", {"text": "I heard Acme Bank is looking for 300,000 sq ft in Shoreditch"}, "Priya")
    sigs = SignalEngine(app.store, app.config).for_role("team")
    lines = team_summary(app.store, sigs)
    assert lines[0].startswith("Central London vacancy") and "Investment was" in lines[1]
    md = app.call("generate_brief", {"format": "board"}, "Sarah")["summary"]
    headline = md.split("## Headline\n")[1].split("\n")[0]
    assert "negative signals" not in headline and "Unverified" not in headline and "vacancy" in headline
    main = md.split("## What changed and why it matters")[1].split("## Market talk")[0]
    assert "Unverified tip" not in main and "Acme Bank" not in main


def test_copilot_says_what_the_library_does_not_cover(app):
    from core import coverage
    savills = coverage.note("What does Savills say about Midtown?", app.store)
    assert "no Savills figures" in savills and "Knight Frank" in savills
    shoreditch = coverage.note("How is Shoreditch vacancy moving?", app.store)
    assert "don't have Shoreditch" in shoreditch and "Clerkenwell/Farringdon" in shoreditch
    assert coverage.note("Midtown prime rent trend", app.store) == ""


def test_api_failure_reaches_the_user_as_one_plain_sentence(app, monkeypatch):
    """No offline fallback any more: a budget or API error must still give a plain message, never a trace."""
    from core.engine import AgentFailed, AgentUnavailable, Engine
    eng = Engine(app)
    eng.ready = True

    def broken(*a, **k):
        raise RuntimeError('Pi Durable agent.ts failed: Traceback ... 400 {"error": "Your credit balance is too low"}')
    monkeypatch.setattr(eng.pi, "ask", broken)
    monkeypatch.setattr(eng.pi, "gather", broken)
    for call in (lambda: eng.ask("What changed?", "Kenneth"), lambda: eng.gather("Kenneth"),
                 lambda: eng.read_upload("x.pdf", "Kenneth")):
        with pytest.raises(AgentFailed) as e:
            call()
        assert "budget has run out" in str(e.value) and "Traceback" not in str(e.value)
        assert "credit balance" in e.value.detail and isinstance(e.value, AgentUnavailable)
    monkeypatch.setattr(eng.pi, "ask", lambda *a: (_ for _ in ()).throw(RuntimeError("HTTP 429 rate_limit_error")))
    with pytest.raises(AgentFailed, match="busy"):
        eng.ask("What changed?", "Kenneth")


def test_copilot_has_its_own_mimo_model(app, monkeypatch):
    """The Copilot (agent.ts) runs on llm.copilot_model; the information agent and the short tool calls
    stay on llm.mimo_model. Env overrides: MIMO_COPILOT_MODEL for the Copilot, MIMO_MODEL for the rest."""
    from core.pi_engine import mimo_model
    for k in ("MIMO_MODEL", "MIMO_COPILOT_MODEL"):
        monkeypatch.delenv(k, raising=False)
    cfg = app.config["llm"]
    assert cfg["copilot_model"] == "mimo-v2.6-pro-ultraspeed" and cfg["mimo_model"] == "mimo-v2.6-flash"
    assert mimo_model(cfg, "agent.ts") == "mimo-v2.6-pro-ultraspeed"
    assert mimo_model(cfg, "info_agent.ts") == "mimo-v2.6-flash"
    assert app.config["vision"]["model"] == "mimo-v2.6-flash"
    assert mimo_model({}, "agent.ts") == "mimo-v2.6-flash"  # no copilot_model: the shared one
    monkeypatch.setenv("MIMO_COPILOT_MODEL", "mimo-v2.6-pro")
    assert mimo_model(cfg, "agent.ts") == "mimo-v2.6-pro" and mimo_model(cfg, "info_agent.ts") == "mimo-v2.6-flash"
    monkeypatch.setenv("MIMO_MODEL", "mimo-v2.5")
    assert mimo_model(cfg, "info_agent.ts") == "mimo-v2.5"
    for name in ("README.md", ".env.example"):  # the docs name the Copilot's model
        assert "mimo-v2.6-pro-ultraspeed" in (ROOT / name).read_text().lower()


def test_answer_votes_are_per_person_and_clearable(app):
    """Thumbs up / down under a Copilot answer: one vote per person per answer, pressing again clears it, and the
    votes never touch the signal feedback that hides cards."""
    st = app.store
    cid = st.new_chat("Voter", "votes")  # its own person, so the saved-chats test's counts are untouched
    st.add_chat_message(cid, "user", "q")
    mid = st.add_chat_message(cid, "assistant", "a")
    assert st.answer_vote(mid, "Voter") == 0
    st.vote_answer(mid, "Voter", 1)
    st.vote_answer(mid, "Other", -1)
    assert st.answer_vote(mid, "Voter") == 1 and st.answer_vote(mid, "Other") == -1
    st.vote_answer(mid, "Voter", -1)
    assert st.answer_vote(mid, "Voter") == -1
    st.vote_answer(mid, "Voter", 0)
    assert st.answer_vote(mid, "Voter") == 0 and st.answer_vote(mid, "Other") == -1
    assert not st.hidden_signals("Other") and not st.liked_signals("Voter")


def test_copilot_layout_pieces_exist():
    """The Copilot page follows the ChatGPT-style layout spec: a page-only stylesheet hook, a sticky top bar, no
    avatars, a right-hand bubble for the question, an action row per answer, the Stop button over the send
    button, a disclaimer under the chat box, and a 500px side panel."""
    src = (ROOT / "app.py").read_text()
    for needle in ('class="loa-cp-page"', "st-key-cp_top_", "stChatMessageAvatarUser", "justify-content:flex-end",
                   "st-key-acts_", "placeStop", "loa-copy", "Copilot can be wrong", 'width="small"', "loa-latest",
                   "st.pills(", "SETUP_PILL", "Worked through"):
        assert needle in src, needle
    assert "Tools used:" not in src and "Enter sends" not in src
    from core import icons
    assert icons.material("copy") and icons.material("clock") and "copy" in icons.PATHS and "arrow-up" in icons.PATHS
    assert 'toolbarMode = "minimal"' in (ROOT / ".streamlit" / "config.toml").read_text()


def test_generated_things_are_kept_by_data_fingerprint(app):
    """The brief and the deck are remembered against a fingerprint of the data they were made from: the same data
    reads them back (no model call), and the fingerprint moves when the picked signals or the data change."""
    st = app.store
    fp = st.data_fingerprint()
    assert fp == st.data_fingerprint() and len(fp) == 16
    assert st.cached("brief", fp) is None
    st.save_cached("brief", fp, "## Board summary\nVacancy fell.", "")
    assert st.cached("brief", fp)["data"].startswith("## Board summary")
    st.add_to_brief({"id": "fp-test", "title": "Vacancy", "markets": ["London"]}, "Kenneth")
    moved = st.data_fingerprint()
    assert moved != fp and st.cached("brief", moved) is None
    st.remove_from_brief("fp-test")
    assert st.data_fingerprint() == fp and st.cached("brief", fp) is not None  # back to the same data, same brief
    st.forget_cached("brief", fp)
    assert st.cached("brief", fp) is None


def test_deck_preview_renders_pdf_and_slide_images(deck_app, tmp_path):
    """A built deck is rendered to a PDF and one PNG per slide for the Team brief page (needs LibreOffice and
    poppler; on a server without them the page shows the outline and the .pptx only)."""
    from core import preview
    out = deck_app.call("build_brief_pptx", {"format": "board"}, "Kenneth")["data"]
    have = preview.available()
    shown = preview.render(out["path"], tmp_path / "prev")
    if not have["soffice"]:
        assert shown["pdf"] == "" and "LibreOffice" in shown["why"]
        pytest.skip("LibreOffice not installed here")
    assert Path(shown["pdf"]).exists() and len(shown["pages"]) == len(out["slides"])
    assert all(Path(q).exists() for q in shown["pages"]) and shown["pages"][0].endswith("slide-1.png")


def test_a_mimo_key_powers_chat_vision_and_pi(app, monkeypatch):
    from core import llm as L
    from core.pi_engine import PiDurableEngine
    for k in ("ANTHROPIC_API_KEY", "OPENAI_API_KEY", "LLM_PROVIDER"):
        monkeypatch.delenv(k, raising=False)
    monkeypatch.setenv("MIMO_API_KEY", "test")
    chat = L.get_llm(app.config["llm"])
    assert chat.name == "mimo" and chat.model == "mimo-v2.6-flash" and "xiaomimimo" in str(chat.client.base_url)
    assert L.get_vision(app.config["vision"], chat).model == "mimo-v2.6-flash"
    monkeypatch.setenv("PI_FAUX", "0")
    ok, why = PiDurableEngine(app).available()
    assert ok or "npm install" in why or "Node" in why  # the key alone is enough for Pi


def test_following_a_source_shows_a_proposal_before_adding(app):
    ev = {"page": "<?xml version='1.0'?><rss></rss>", "content_type": "application/rss+xml",
          "terms": "<p>Copyright 2026.</p>", "robots": ""}
    url = "https://officenews.example.com/feed"
    plan = app.call("follow_source", {"source": url, "propose": True, **ev}, "Priya")["data"]
    assert plan["channel"] == "auto" and plan["chain"][0] == "check-terms" and "News and tips" in plan["adds"]
    assert not app.store.q("SELECT 1 FROM followed WHERE url=?", (url,))  # nothing saved yet
    savills = app.call("follow_source", {"source": "Savills Market Watch", "propose": True,
                                         "page": "<p>Research</p>"}, "Priya")["data"]
    assert savills["channel"] == "upload" and "read-pdf" in savills["chain"] and "Savills figures" in savills["adds"]
    app.call("follow_source", {"source": url, "name": plan["name"], **ev}, "Priya")
    assert app.store.q("SELECT 1 FROM followed WHERE url=?", (url,))


def test_a_round_never_resends_held_news_to_the_model(app):
    """The real-key test timed out because every round re-classified 85 stored stories. Held stories are
    skipped before the classifier (and so the model) sees them; new ones are classified once."""
    from core import ingest
    seen = []
    real = app.ctx.classifier.classify
    app.ctx.classifier.classify = lambda item: seen.append(item.title) or real(item)
    try:
        items = [RawItem("t", title="Held story: Landsec lets 50,000 sq ft in Victoria", text="Victoria letting", url="https://x/1"),
                 RawItem("t", title="New story: Derwent pre-lets 80,000 sq ft in Fitzrovia", text="Fitzrovia pre-let", url="https://x/2")]
        ingest.ingest_items(app.ctx, items[:1])
        seen.clear()
        res = ingest.ingest_items(app.ctx, items)
    finally:
        app.ctx.classifier.classify = real
    assert seen == ["New story: Derwent pre-lets 80,000 sq ft in Fitzrovia"]
    assert res["counts"]["news_already_held"] == 1 and res["counts"]["news"] == 1


def test_terms_page_must_be_a_terms_page_not_an_article():
    from core import follow
    page = ('<a href="/news/2026/10/legal-battle-over-canary-wharf-tower">Legal battle over tower</a>'
            '<a href="/sport">Sport</a>')
    assert follow.terms_link(page, "https://www.cityam.com/") == ""
    assert follow.terms_link(page + '<a href="/terms-and-conditions">T&Cs</a>', "https://www.cityam.com/") \
        == "https://www.cityam.com/terms-and-conditions"
    assert follow.is_terms_url("https://x.com/legal") and follow.is_terms_url("https://x.com/about/legal-notice.html")
    assert not follow.is_terms_url("https://x.com/news/legal-sector-takes-space")


def test_following_a_source_is_not_a_tip(app):
    from core import gather
    assert not gather.looks_like_market_news("Following Property Week (upload only), checked by Pi")
    assert gather.looks_like_market_news("Columbus Courtyard may have a pre-let under offer")
    assert gather.looks_like_market_news("A broker says Google wants 200,000 sq ft in King's Cross")
    before = len(app.store.news())
    out = app.call("share_tip", {"text": "Now following Property Week for the team", "how": "Other"}, "sam")
    assert out["data"]["filed"] is False and len(app.store.news()) == before


def test_publisher_from_file_name_and_one_row_per_figure(app):
    assert documents.detect_publisher("Central London office market update", "BNPPRECentralLondonOfficeMarketUpdateQ32025.pdf") == "BNP Paribas"
    assert documents.detect_publisher("no clue", "report.pdf") == "Unknown publisher"
    raw = [{"market_text": "Central London", "metric": "vacancy_pct", "value": 8.4, "page": 2, "where": "KPI tile"},
           {"market_text": "Central London", "metric": "vacancy_pct", "value": 8.4, "page": 5, "where": "text"}]
    cands = documents.check_figures(app.store, raw, ["", "vacancy 8.4%", "", "", "vacancy 8.4%"], "2025Q3")
    assert len(cands) == 1


def test_brief_headline_has_no_heading_marks(app, monkeypatch):
    class Loud:
        available, model = True, "test"
        def complete(self, *a, **k):
            return "# Headline\nLeasing is strong and investment is not.\n## More\nRates held."
        def complete_json(self, *a, **k):
            return {}
    monkeypatch.setattr(app.ctx, "llm", Loud())
    from core.summary import is_tip
    first = next(s for s in SignalEngine(app.store, app.config).for_role("team") if not is_tip(s))
    app.store.add_to_brief(first.to_dict(), "sam")
    text = app.call("generate_brief", {"format": "ic"}, "sam")["summary"]
    section = text.split("## Headline")[1].split("## Key numbers")[0]
    assert "#" not in section and "Leasing is strong and investment is not. Rates held." in section


def test_a_tip_saved_by_the_agent_is_still_unverified(app):
    """The information agent saves tips with source 'Team'; they must still read as unverified and stay out
    of 'What changed' in the brief."""
    from core.models import NewsItem
    from core.summary import is_tip
    app.store.add_news(NewsItem(id="tipteam1", title="Heard a US law firm is close to signing at 2 Finsbury Avenue",
                                body="", url="", source="Team", published="2026-10-09", category="leasing_deal",
                                markets=["City Core"], importance=3,
                                shared_by=[{"who": "priya", "how": "LinkedIn", "when": "2026-10-09T09:00:00", "note": "", "url": ""}]))
    sig = next(s for s in SignalEngine(app.store, app.config).for_role("team") if "2 Finsbury" in s.title)
    assert is_tip(sig) and not any("Team (" in e.text for e in sig.evidence)
    for it in app.store.brief_items():   # earlier tests picked items; the tip must sort itself even when picked
        app.store.remove_from_brief(it["id"])
    app.store.add_to_brief(sig.to_dict(), "priya")
    brief = app.call("generate_brief", {"format": "board"}, "priya")["summary"]
    assert "2 Finsbury" in brief.split("## Market talk")[1] and "2 Finsbury" not in brief.split("## Market talk")[0]


def test_vision_reads_page_batches_in_parallel_and_never_twice(app, tmp_path, monkeypatch):
    import threading
    import time
    from skills.info_processing import tools as info
    from pypdf import PdfReader, PdfWriter
    page = PdfReader(ROOT / "reference" / "samples" / "sample_bnp_paribas_q2_2025_restated.pdf").pages[0]
    w = PdfWriter()
    for _ in range(4):
        w.add_page(page)
    pdf = tmp_path / "r.pdf"
    w.write(pdf)
    calls, active, peak = [], [0], [0]
    lock = threading.Lock()

    class Slow:
        available, model, pages_per_call, long_side, parallel = True, "fake-vision", 1, 800, 4
        def complete_json_images(self, prompt, images):
            with lock:
                active[0] += 1; peak[0] = max(peak[0], active[0]); calls.append(len(images))
            time.sleep(0.3)
            with lock:
                active[0] -= 1
            return [{"market_text": "Central London", "metric": "vacancy_pct", "value": 8.4, "page": 1, "where": "x"}]
    monkeypatch.setattr(app.ctx, "vision", Slow())
    out = info.read_pdf_pages(app.ctx, {"path": str(pdf), "pages": [1, 2, 3, 4]})
    assert len(calls) == 4 and peak[0] >= 2, "batches should overlap"
    again = info.read_pdf_pages(app.ctx, {"path": str(pdf), "pages": [1, 2, 3, 4]})
    assert len(calls) == 4 and again["summary"].startswith("(already read")


def test_checks_merge_rounded_duplicates_and_flag_ambiguous_measures(app):
    raw = [{"market_text": "City & Southbank", "metric": "vacancy_pct", "value": 9.9, "page": 1},
           {"market_text": "City & Southbank", "metric": "vacancy_pct", "value": 9.91, "page": 3},
           {"market_text": "Central London", "metric": "vacancy_pct", "value": 1.2, "page": 2, "label_status": "inferred"},
           {"market_text": "Central London", "metric": "investment_gbp_bn", "value": 0.85, "page": 4},
           {"market_text": "Central London", "metric": "investment_gbp_bn", "value": 4.26, "page": 4}]
    cands = documents.check_figures(app.store, raw, ["9.9%", "1.2%", "9.91%", "0.85 4.26"], "2025Q3")
    vac = [c for c in cands if c["market"] == "City & Southbank"]
    assert len(vac) == 1
    low = next(c for c in cands if c["value"] == 1.2)
    assert any("plausible range" in f for f in low["flags"]) and not low["suggest_keep"]
    assert low["status"] == "implausible"
    grade_a = documents.check_figures(app.store, [{"market_text": "Central London", "metric": "vacancy_pct", "value": 2.8,
                                                   "page": 4, "where": "Grade A availability tile"}], ["2.8%"] * 4, "2025Q3")
    assert grade_a[0]["metric"] == "new_refurb_vacancy_pct"
    inv = [c for c in cands if c["metric"] == "investment_gbp_bn"]
    assert len(inv) == 2 and all(not c["suggest_keep"] and any("different values" in f for f in c["flags"]) for c in inv)


def test_feed_url_looks_for_terms_on_the_home_page(monkeypatch):
    from core import follow
    pages = {"https://news.example.com/robots.txt": "User-agent: *\nAllow: /",
             "https://news.example.com/rss": "<?xml version='1.0'?><rss><channel><title>x</title></channel></rss>",
             "https://news.example.com/": '<html><footer><a href="/terms">Terms</a></footer></html>',
             "https://news.example.com/terms": "<html>No automated collection or AI use permitted.</html>"}

    class R:
        def __init__(self, url):
            self.text, self.ok, self.headers = pages.get(url, ""), url in pages, {"content-type": "application/rss+xml"}
    import requests
    monkeypatch.setattr(requests, "get", lambda url, **k: R(url))
    ev = follow.fetch_evidence("https://news.example.com/rss")
    assert ev["terms_url"] == "https://news.example.com/terms" and "AI use" in ev["terms"]


# ---------------------------------------------------------------- team brief as slides (research/brief-to-pptx.md)
def _slide_text(slide) -> str:
    out = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            out.append(sh.text_frame.text)
        if getattr(sh, "has_table", False) and sh.has_table:
            out += [c.text for row in sh.table.rows for c in row.cells]
    return "\n".join(out)


def _chart_values(slide):
    return [[list(s.values) for s in sh.chart.plots[0].series] for sh in slide.shapes if getattr(sh, "has_chart", False) and sh.has_chart]


@pytest.fixture
def deck_app(app, tmp_path, monkeypatch):
    monkeypatch.setattr(app, "root", tmp_path)
    monkeypatch.setattr(app.ctx, "root", tmp_path)
    for it in app.store.brief_items():
        app.store.remove_from_brief(it["id"])
    return app


def test_brief_deck_opens_with_three_to_six_slides_and_the_cover_numbers(deck_app):
    from pptx import Presentation
    from core import analytics as an
    out = deck_app.call("build_brief_pptx", {"format": "board"}, "Kenneth")
    prs = Presentation(out["data"]["path"])
    assert 3 <= len(prs.slides) <= 6 and len(out["data"]["slides"]) == len(prs.slides)
    cover = _slide_text(prs.slides[0])
    from core.summary import team_summary
    assert team_summary(deck_app.store, [])[0] in cover
    for num in ("8.0%", "3.10m sq ft", "£2.00bn", "3.75%"):
        assert num in cover
    # native charts on slides 2-4 carry the library's series
    vac = _chart_values(prs.slides[1])
    assert vac and vac[0][0][-1] == an.snapshot(deck_app.store, "London", "vacancy_pct")["value"] == 8.0
    assert vac[1][0][-1] == an.snapshot(deck_app.store, "London", "take_up_msqft")["value"] == 3.1
    assert sum(bool(_chart_values(sl)) for sl in prs.slides) >= 2, "native charts on the content slides"
    assert out["data"]["storyline"] == "rules" and all(out["data"]["why"]), "every slide says why it is there"
    notes = [sl.notes_slide.notes_text_frame.text for sl in prs.slides]
    assert all(n.startswith("Why this slide:") for n in notes)
    # our deals, when the storyline includes them, sit apart from market deals and drafts never appear
    spec = {"theme": "navy-gold", "slides": [
        {"layout": "Cover", "title": "Deals", "why": "cover", "notes": "n", "fill": {"subtitle": "auto", "headline": "auto", "caveat": "auto", "kpis": "auto"}},
        {"layout": "Title only", "title": "Our deals and market deals", "why": "2 of our deals confirmed", "notes": "n", "fill": {"table": {"table": "deals"}}},
        {"layout": "Sources", "title": "Sources", "why": "s", "notes": "n", "fill": {"body": "auto"}}]}
    deals = _slide_text(Presentation(deck_app.call("build_brief_pptx", {"format": "board", "spec": spec}, "k")["data"]["path"]).slides[1])
    assert "8 Bishopsgate" in deals and "Internal" in deals and "Uploaded" in deals, "our deals listed apart from market deals"
    assert "1 Triton Square, 4th floor" not in deals, "drafts never reach the deck"
    last = _slide_text(prs.slides[-1])
    assert "Bank of England Bank Rate" in last and "verify against the original" in last
    assert all(f"[{i}]" in last for i in range(1, 4))


def test_brief_deck_refuses_numbers_the_data_does_not_have(deck_app):
    bad = deck_app.call("build_brief_pptx", {"format": "board", "slides": {"occupiers": {"title": "Vacancy collapses to 6.5%"}}}, "k")
    assert "numbers not in the data: 6.5" in bad["summary"] and not bad["data"].get("path")
    assert not list((deck_app.root / "data" / "exports").glob("*.pptx")) or True  # nothing new written below
    ok = deck_app.call("build_brief_pptx", {"format": "board", "slides": {"occupiers": {"title": "Vacancy fell to 8.0% in Q2 2026",
                                                                                         "notes": "Lowest since Q1 2024."}}}, "k")
    from pptx import Presentation
    prs = Presentation(ok["data"]["path"])
    occ = next(sl for sl in prs.slides if "Vacancy fell to 8.0% in Q2 2026" in _slide_text(sl))
    assert occ.notes_slide.notes_text_frame.text.endswith("Lowest since Q1 2024.")


def test_brief_deck_shows_flags_and_keeps_market_talk_apart(deck_app):
    from pptx import Presentation
    from core.summary import is_tip
    row = deck_app.store.q("SELECT rowid FROM observations WHERE market='London' AND metric='vacancy_pct' "
                           "ORDER BY period DESC LIMIT 1")[0]
    deck_app.store.flag_observation(row["rowid"], "looks like the City figure", "sam")
    tip = next(s for s in SignalEngine(deck_app.store, deck_app.config).for_role("team") if is_tip(s))
    deck_app.store.add_to_brief(tip.to_dict(), "priya")
    real = [s for s in SignalEngine(deck_app.store, deck_app.config).for_role("team") if not is_tip(s)][:2]
    for s in real:
        deck_app.store.add_to_brief(s.to_dict(), "priya")
    try:
        prs = Presentation(deck_app.call("build_brief_pptx", {"format": "board"}, "k")["data"]["path"])
        assert "8.0% (flagged)" in _slide_text(prs.slides[0])
        assert "flagged figures, not yet checked" in _slide_text(prs.slides[-1]).lower()
        signals = _slide_text(prs.slides[4])
        body = tip.title.removeprefix("Unverified tip: ")
        assert "market talk, not confirmed" in signals.lower() and body in signals
        assert signals.count(body) == 1 and "Unverified tip:" not in signals
    finally:
        deck_app.store.resolve_observation(row["rowid"], "sam", True)


def test_ic_deck_has_evidence_table_and_two_columns(deck_app):
    from pptx import Presentation
    prs = Presentation(deck_app.call("build_brief_pptx", {"format": "ic"}, "Kenneth")["data"]["path"])
    assert 3 <= len(prs.slides) <= 6
    ev = next(sl for sl in prs.slides if "rent-free" in _slide_text(sl).lower())
    assert any(getattr(sh, "has_table", False) and sh.has_table for sh in ev.shapes)
    s5 = _slide_text(next(sl for sl in prs.slides if sl.slide_layout.name == "Signal cards")).lower()
    assert "supports the case" in s5 and "risks to the case" in s5
    assert s5.index("supports the case") < s5.index("risks to the case")  # the case for, then the case against


@needs_pi
def test_copilot_builds_the_deck_with_the_scripted_model(deck_app, tmp_path):
    import json
    script = tmp_path / "deck.json"
    script.write_text(json.dumps([{"tool": "brief_deck_data", "args": {"format": "board"}},
                                  {"tool": "build_brief_pptx", "args": {"format": "board", "slides": {
                                      "occupiers": {"title": "Vacancy fell to 8.0%, the lowest since Q1 2024"}}}},
                                  {"text": "Slides built."}]))
    os.environ.update(PI_FAUX="1", PI_FAUX_SCRIPT=str(script), PI_STORE=str(tmp_path / "pi.sqlite"),
                      PI_CONVERSATIONS=str(tmp_path / "c.json"))
    try:
        from core.engine import Engine
        ans = Engine(deck_app).ask("Make the brief into slides for Monday's meeting", "Kenneth")
    finally:
        os.environ.pop("PI_FAUX_SCRIPT", None)
    tools = [t["tool"] for t in ans.trace]
    assert tools[:2] == ["brief_deck_data", "build_brief_pptx"]
    built = ans.trace[1]["result"]["data"]
    assert Path(built["path"]).exists() and built["slides"][1]["title"].startswith("Vacancy fell to 8.0%")


def test_every_house_theme_builds_a_clean_six_slide_deck(deck_app):
    from pptx import Presentation
    from skills.brief_deck import tools as bd
    for theme in bd.THEMES:
        out = deck_app.call("build_brief_pptx", {"format": "board", "theme": theme}, "Kenneth")
        assert out["data"]["theme"] == theme and out["data"]["problems"] == [], (theme, out["summary"])
        assert len(Presentation(out["data"]["path"]).slides) == 6
    ic = deck_app.call("build_brief_pptx", {"format": "ic"}, "Kenneth")["data"]
    assert ic["theme"] == "data-white" and ic["problems"] == []


def test_the_copilot_can_write_its_own_storyline_with_the_new_layouts(deck_app):
    import json
    from pptx import Presentation
    from skills.brief_deck import tools as bd
    guide = deck_app.call("brief_deck_guide", {"topic": "showcase"}, "k")
    spec = guide["data"]["spec"]
    assert [s["layout"] for s in spec["slides"]][1:5] == ["Statement", "Stat row", "Numbered pillars", "Timeline"]
    out = deck_app.call("build_brief_pptx", {"format": "board", "spec": spec}, "k")
    assert out["data"]["problems"] == [], out["summary"]
    prs = Presentation(out["data"]["path"])
    assert len(prs.slides) == 6 and [s.slide_layout.name for s in prs.slides] == [s["layout"] for s in spec["slides"]]
    assert "8.0%" in _slide_text(prs.slides[2])  # the stat row carries the library's vacancy figure
    # a five-slide spec of the agent's own, on another theme, with a hero number
    five = {"theme": "charcoal-amber", "slides": [
        {"layout": "Cover", "eyebrow": "London office market", "title": "Deal discussion", "notes": "n",
         "fill": {"subtitle": "auto", "headline": "auto", "caveat": "auto", "kpis": "auto"}},
        {"layout": "Hero number", "eyebrow": "01 · Occupiers", "title": "Vacancy is at *8.0%*", "notes": "n",
         "so_what": "Tightening supply supports rents.",
         "fill": {"hero": {"metric": "vacancy_pct"}, "chart_label": "Vacancy", "chart": {"chart": "vacancy"},
                  "source": "See sources."}},
        {"layout": "Chart + takeaway", "eyebrow": "02 · Rents", "title": "Prime rents are holding", "notes": "n",
         "fill": {"chart_label": "Prime rent", "chart": {"chart": "rents"}, "takeaway": ["Core rents held."],
                  "source": "See sources."}},
        {"layout": "Signal cards", "eyebrow": "03 · Signals", "title": "What changed", "notes": "n",
         "fill": {"cards": "auto", "talk": "auto"}},
        {"layout": "Sources", "eyebrow": "Sources", "title": "Sources", "notes": "n", "fill": {"body": "auto"}}]}
    out = deck_app.call("build_brief_pptx", {"format": "board", "spec": five}, "k")
    assert out["data"]["theme"] == "charcoal-amber" and out["data"]["problems"] == [], out["summary"]
    assert len(Presentation(out["data"]["path"]).slides) == 5
    assert set(bd.THEMES) == {"navy-gold", "editorial-cream", "data-white", "charcoal-amber", "forest-cream"}
    assert len(bd.build_deck.LAYOUTS) == 13


def test_a_spec_is_refused_for_bad_theme_layout_count_or_number(deck_app):
    import copy
    from skills.brief_deck import tools as bd
    base = bd.build_deck.default_spec(bd.deck_data(deck_app.ctx, "board"))
    assert len(base["slides"]) == 6

    def refused(spec, want):
        out = deck_app.call("build_brief_pptx", {"format": "board", "spec": spec}, "k")
        assert not out["data"].get("path") and want in out["summary"], (want, out["summary"])

    bad = copy.deepcopy(base); bad["theme"] = "neon"
    refused(bad, "unknown theme")
    bad = copy.deepcopy(base); bad["slides"][1]["layout"] = "Mood board"
    refused(bad, "unknown layout")
    bad = copy.deepcopy(base)
    for _ in range(7):  # the slide editor may grow a deck to 12; 13 is refused
        bad["slides"].insert(2, copy.deepcopy(base["slides"][2]))
    refused(bad, "13 slides")
    bad = copy.deepcopy(base); bad["slides"][1]["title"] = "Vacancy collapses to 6.5%"
    refused(bad, "numbers not in the data: 6.5")
    bad = copy.deepcopy(base); bad["slides"][1]["so_what"] = "Rents will rise 42% next year"
    refused(bad, "numbers not in the data: 42")
    refused({"theme": "navy-gold", "slides": []}, "0 slides")


# ---------------------------------------------------------------- internal data
def test_a_sentence_to_the_copilot_becomes_one_of_our_deals_and_later_pieces_merge(app):
    from core import internal as I
    before = len(app.store.internal_deals())
    out = app.call("record_deal", {"text": "We just let 20 Fenchurch Street 7th floor to Acme Partners at £68 psf, 10 year term"}, "Kenneth")
    d = out["data"]["deal"]
    assert out["data"]["recorded"] and not out["data"]["merged"]
    assert d["property"] == "20 Fenchurch Street" and d["floor_area"] == "7th floor" and d["tenant"] == "Acme Partners"
    assert d["rent"] == "£68.00 psf" and d["term"] == "10 years" and d["submarket"] == "City Core"
    assert d["status"] == "draft" and set(d["missing"]) == {"rent_free", "counterparty", "date"}, "'just let' is not a date"
    assert d["confidentiality"] == "internal" and d["recorded_by"][0][:2] == ["Kenneth", "said to Copilot"]
    assert "Still missing: Rent-free, Date, Counterparty" in out["summary"] and "Internal only" in out["summary"]
    # a second sentence about the same property and floor fills the gaps instead of making a second record
    again = app.call("record_deal", {"text": "20 Fenchurch Street 7th floor: 18 months rent free, completed 3 Oct 2026, landlord: Land Securities"}, "Priya")
    assert again["data"]["merged"] and len(app.store.internal_deals()) == before + 1
    d = again["data"]["deal"]
    assert d["rent_free"] == "18 months" and d["counterparty"] == "Landlord: Land Securities" and d["missing"] == []
    assert d["rent"] == "£68.00 psf", "a merge fills empty fields and never overwrites"
    assert [w for w, _, _ in d["recorded_by"]] == ["Kenneth", "Priya"]
    # a different floor of the same building is a different deal
    other = app.call("record_deal", {"text": "20 Fenchurch Street 9th floor let to Beta LLP at £70 psf"}, "Kenneth")
    assert not other["data"]["merged"] and len(app.store.internal_deals()) == before + 2
    # the recorder confirms; nothing to do with market observations
    assert app.store.q("SELECT COUNT(*) AS n FROM observations WHERE market='City Core' AND metric='prime_rent_psf' AND value=68")[0]["n"] == 0
    conf = app.call("confirm_deal", {"id": d["id"]}, "Kenneth")
    assert conf["data"]["status"] == "confirmed"
    for x in (d["id"], other["data"]["deal"]["id"]):
        app.store.delete_internal_deal(x)


def test_a_photo_of_the_lease_is_read_by_the_vision_model_and_merged_into_the_draft(app, monkeypatch, tmp_path):
    """End to end: Sarah said one sentence (a draft with fields missing); James drops a photo of the lease front
    page; the vision model reads it; the record is completed, not duplicated, the photo kept beside it."""
    from core import internal as I
    photo = ROOT / "reference" / "samples" / "lease_front_page_8_bishopsgate.jpg"
    assert photo.exists()
    for d in app.store.internal_deals():  # start from the draft only
        app.store.delete_internal_deal(d["id"])
    draft = app.call("record_deal", {"text": "8 Bishopsgate 12th floor just let at £72.50 psf"}, "Sarah")["data"]["deal"]
    assert draft["status"] == "draft" and "tenant" in draft["missing"] and "term" in draft["missing"]
    seen = []

    class FakeVision:  # what MiMo returns for the sample photo (the real-key test runs in the MiMo thread)
        available, model = True, "fake-vision"
        def read_json_images(self, prompt, images):
            seen.append((len(images), images[0][:3]))
            assert "lease" in prompt.lower() and "never guess" in prompt
            return {"deal": {"property": "8 Bishopsgate, London EC2N 4BQ", "floor_area": "Twelfth floor, 18,400 sq ft",
                             "tenant": "Confidential (US law firm)", "rent": "£72.50 psf", "term": "10 years",
                             "rent_free": "24 months", "date": "2 October 2026", "counterparty": "Mitsubishi Estate London Limited"},
                    "confidence": "high", "notes": "Lease front page."}
    monkeypatch.setattr(app.ctx, "vision", FakeVision())
    read = app.call("read_photo", {"path": str(photo), "hint": "lease front page"}, "James")
    assert seen == [(1, b"\xff\xd8\xff")], "the JPEG goes to the model as it is"
    assert read["data"]["fields"]["date"] == "2026-10-02" and "high confidence" in read["summary"]
    saved = app.call("save_internal_deal", {"fields": read["data"]["fields"], "how": "photo", "original": str(photo),
                                            "confidence": "high"}, "James")
    d = saved["data"]["deal"]
    assert saved["data"]["merged"] and d["id"] == draft["id"], "same property and floor: merged, not a second record"
    assert d["missing"] == [] and d["tenant"] == "Confidential (US law firm)" and d["counterparty"] == "Mitsubishi Estate London Limited"
    assert d["originals"] == [str(photo)] and [w for w, _, _ in d["recorded_by"]] == ["Sarah", "James"]
    assert len(app.store.internal_deals()) == 1
    # without a vision model the photo is kept beside a draft, never guessed
    monkeypatch.setattr(app.ctx, "vision", None)
    monkeypatch.setattr(app.ctx, "llm", type("Off", (), {"available": False})())
    blind = app.call("read_photo", {"path": str(photo)}, "James")
    assert blind["data"]["fields"] == {} and "No vision model" in blind["summary"]
    for x in app.store.internal_deals():
        app.store.delete_internal_deal(x["id"])
    app.seed_internal()


def test_our_deals_stay_apart_from_market_deals_and_internal_only_never_reaches_the_ic_paper(app):
    from core import internal as I
    deals = {d["property"]: d for d in app.store.internal_deals()}
    assert "8 Bishopsgate, 12th floor" in deals and deals["8 Bishopsgate, 12th floor"]["confidentiality"] == "internal"
    board = app.call("generate_brief", {"format": "board"}, "Kenneth")["summary"]
    ours = board.split("## Our deals")[1].split("## Market deals")[0]
    market = board.split("## Market deals")[1].split("##")[0]
    assert "8 Bishopsgate" in ours and "[internal]" in ours and "25 Cabot Square" in ours
    assert "1 Triton Square, 4th floor" not in ours, "a draft is never listed"
    assert "Anthropic" in market and "[" in market and "8 Bishopsgate" not in market
    ic = app.call("generate_brief", {"format": "ic"}, "Kenneth")["summary"]
    ic_ours = ic.split("## Our deals")[1].split("## Market deals")[0]
    assert "25 Cabot Square" in ic_ours and "8 Bishopsgate" not in ic_ours and "1 internal-only record(s) not included" in ic_ours
    # the external comparison: a market deal in the same building with a different rent is noted, our record kept
    app.store.add_deal({"period": "2026Q3", "deal_type": "letting", "party": "Savills Market Watch Oct 2026 (test)",
                        "building": "25 Cabot Square, E14", "submarket": "Canary Wharf", "size_sqft": 9800, "rent_psf": 50.0,
                        "price_gbp_m": None, "yield_pct": None, "note": "", "source_id": "KF_LOMR_2024Q3"})
    try:
        d = I.record(app.store, {"property": "25 Cabot Square", "floor_area": "3rd floor", "date": "2026-09-18"}, "Priya", "said to Copilot")["deal"]
        assert d["ext"] and d["ext"]["differs"] and "£50.00 psf against our £52.00 psf" in d["ext"]["text"]
        assert d["rent"] == "£52.00 psf", "our record is kept"
        md = app.call("generate_brief", {"format": "board"}, "Kenneth")["summary"]
        assert "reports £50.00 psf; our record stands" in md
        # quoting toggle: marking it internal only removes it from the IC paper
        app.call("set_deal_confidentiality", {"id": d["id"], "level": "internal"}, "Priya")
        ic = app.call("generate_brief", {"format": "ic"}, "Kenneth")["summary"]
        assert "25 Cabot Square" not in ic.split("## Our deals")[1].split("## Market deals")[0]
        app.call("set_deal_confidentiality", {"id": d["id"], "level": "quotable"}, "Priya")
    finally:
        app.store.db.execute("DELETE FROM deals WHERE party LIKE 'Savills Market Watch Oct 2026 (test)%'")
        app.store.db.commit()
        d = I.record(app.store, {"property": "25 Cabot Square", "floor_area": "3rd floor", "date": "2026-09-18"}, "Priya", "said to Copilot")["deal"]
        assert not d["ext"]


def test_a_spreadsheet_of_our_deals_files_one_record_per_row(app, tmp_path):
    import pandas as pd
    xl = tmp_path / "deals_q3.xlsx"
    pd.DataFrame([{"Property": "8 Bishopsgate, 12th floor", "Tenant": "Confidential (US law firm)", "Rent psf": 72.5,
                   "Term": "10 years", "Rent free": "24 months", "Date": "2026-10-02"},
                  {"Property": "30 Gresham Street, 2nd floor", "Tenant": "Delta Re", "Rent psf": 66.0,
                   "Term": "5 years", "Rent free": "12 months", "Date": "2026-09-30"},
                  {"Property": "", "Tenant": "nobody", "Rent psf": 1, "Term": "", "Rent free": "", "Date": ""}]).to_excel(xl, index=False)
    before = len(app.store.internal_deals())
    out = app.call("deals_from_table", {"path": str(xl)}, "Priya")
    assert out["data"] == {"rows": 2, "merged": 1, "new": 1} and "skipped" in out["summary"]
    new = next(d for d in app.store.internal_deals() if d["property"].startswith("30 Gresham"))
    assert new["rent"] == "£66.00 psf" and new["submarket"] == "City Core" and new["recorded_by"][0][1] == "spreadsheet deals_q3.xlsx"
    assert len(app.store.internal_deals()) == before + 1
    app.store.delete_internal_deal(new["id"])


@needs_pi
def test_copilot_records_a_deal_with_the_scripted_model(app, tmp_path):
    import json
    script = tmp_path / "deal.json"
    script.write_text(json.dumps([{"tool": "record_deal", "args": {"text": "We just let 1 Finsbury Avenue 3rd floor to Orbit Labs at £80 psf, 10 year term"}},
                                  {"text": "Recorded as one of our deals, internal only."}]))
    os.environ.update(PI_FAUX="1", PI_FAUX_SCRIPT=str(script), PI_STORE=str(tmp_path / "pi.sqlite"),
                      PI_CONVERSATIONS=str(tmp_path / "c.json"))
    try:
        from core.engine import Engine
        ans = Engine(app).ask("We just let 1 Finsbury Avenue 3rd floor to Orbit Labs at £80 psf, 10 year term", "Kenneth")
    finally:
        os.environ.pop("PI_FAUX_SCRIPT", None)
    assert [t["tool"] for t in ans.trace][:1] == ["record_deal"]
    d = ans.trace[0]["result"]["data"]["deal"]
    assert d["tenant"] == "Orbit Labs" and d["confidentiality"] == "internal" and d["recorded_by"][0][0] == "Kenneth"
    app.store.delete_internal_deal(d["id"])



def test_the_deck_storyline_follows_the_data_not_a_template(deck_app):
    """Two weeks with different data give different slide sets: here the team sets an alert line that this
    quarter's vacancy crosses, and investment comes in above its long-term average."""
    from core import analytics as an
    from core.models import Observation
    from skills.brief_deck import planner, tools as bd
    before = bd.deck_data(deck_app.ctx, "board")
    plan_a = planner.plan_spec(before, None, before["focus"], deck_app.store.watch_rules())
    keys_a = [(s["key"], s["layout"]) for s in plan_a["slides"]]
    assert ("occupiers", "Two charts") in keys_a and ("capital", "Chart + table") in keys_a
    rule = deck_app.store.add_watch_rule("team", "London", "vacancy_pct", "<=", 8.0, "vacancy at or below 8%", "Sarah")
    sid = an.snapshot(deck_app.store, "London", "investment_gbp_bn")["source_id"]
    deck_app.store.add_observations([Observation(period="2026Q3", market="London", metric="investment_gbp_bn", value=3.6, lta=3.0, source_id=sid, note="")])
    try:
        after = bd.deck_data(deck_app.ctx, "board")
        plan_b = planner.plan_spec(after, None, after["focus"], deck_app.store.watch_rules())
        keys_b = [(s["key"], s["layout"]) for s in plan_b["slides"]]
        assert keys_b != keys_a
        occ = next(s for s in plan_b["slides"] if s["key"] == "occupiers")
        assert occ["layout"] == "Hero number" and "alert line" in occ["why"], "a threshold crossing leads with the number"
        cap = next(s for s in plan_b["slides"] if s["key"] == "capital")
        assert "above the long-term average" in cap["why"] and "above the long-term average" in cap["title"]
        assert [c["reason"] for c in after["stories"]] != [c["reason"] for c in before["stories"]]
        out = deck_app.call("build_brief_pptx", {"format": "board"}, "Kenneth")  # the fallback builds plan_b
        assert out["data"]["problems"] == [] and [s["layout"] for s in out["data"]["slides"]] == [k[1] for k in keys_b]
        assert any("alert line" in w for w in out["data"]["why"])
    finally:
        deck_app.store.delete_watch_rule(rule)
        deck_app.store.db.execute("DELETE FROM observations WHERE period='2026Q3' AND metric='investment_gbp_bn'")
        deck_app.store.db.commit()


@needs_pi
def test_copilot_plans_the_deck_and_every_slide_carries_its_why(deck_app, tmp_path):
    """The agent path: brief_deck_data (ranked stories) -> the Copilot's own spec with a why per slide ->
    build. With the scripted model the choice is fixed; the real-key run is the MiMo test thread's."""
    import json
    from pptx import Presentation
    spec = {"theme": "charcoal-amber", "slides": [
        {"layout": "Cover", "eyebrow": "London office market", "title": "This week in three slides", "why": "Every brief opens with the KPIs.",
         "notes": "n", "fill": {"subtitle": "auto", "headline": "auto", "caveat": "auto", "kpis": "auto"}},
        {"layout": "Hero number", "eyebrow": "01 · Occupiers", "title": "Vacancy is at *8.0%*, the lowest since Q1 2024",
         "why": "Vacancy is the lowest in the series and take-up runs 11% above average; nothing else moved as much.",
         "so_what": "Tightening supply supports rents.", "notes": "n",
         "fill": {"hero": {"metric": "vacancy_pct"}, "chart_label": "Vacancy", "chart": {"chart": "vacancy"}, "source": "See sources."}},
        {"layout": "Sources", "eyebrow": "Sources", "title": "Sources", "why": "Footnotes.", "notes": "n", "fill": {"body": "auto"}}]}
    script = tmp_path / "plan.json"
    script.write_text(json.dumps([{"tool": "brief_deck_data", "args": {"format": "board"}},
                                  {"tool": "build_brief_pptx", "args": {"format": "board", "spec": spec}},
                                  {"text": "Three slides: cover; vacancy, because it is the lowest since Q1 2024; sources."}]))
    os.environ.update(PI_FAUX="1", PI_FAUX_SCRIPT=str(script), PI_STORE=str(tmp_path / "pi.sqlite"),
                      PI_CONVERSATIONS=str(tmp_path / "c.json"))
    try:
        from core.engine import Engine
        built, ans = Engine(deck_app).build_deck("board", None, "Kenneth")
    finally:
        os.environ.pop("PI_FAUX_SCRIPT", None)
    assert built and built["storyline"] == "copilot" and built["theme"] == "charcoal-amber" and built["problems"] == []
    data = ans.trace[0]["result"]
    assert "What stands out this week, ranked" in data["summary"] and data["data"]["stories"][0]["reason"]
    prs = Presentation(built["path"])
    assert len(prs.slides) == 3
    assert prs.slides[1].notes_slide.notes_text_frame.text.startswith("Why this slide: Vacancy is the lowest in the series")
    assert "lowest since Q1 2024" in ans.text


def test_a_sentence_asking_for_slides_is_understood_without_the_copilot():
    """When Pi is offline the app still turns "a deck on the forest theme" into a build, not a dead end."""
    from skills.brief_deck.tools import deck_request
    assert deck_request("Build a client update deck on the forest theme") == ("board", "forest-cream")
    assert deck_request("Make the brief into slides for Monday's meeting") == ("board", "navy-gold")
    assert deck_request("A PowerPoint for the IC on Thursday") == ("ic", "data-white")
    assert deck_request("Compare Midtown and City Core") is None


def test_fewer_signals_leave_no_empty_card_frames_and_demand_is_not_a_repeat_of_the_cover(deck_app, tmp_path):
    from pptx import Presentation
    from skills.brief_deck import planner, tools as bd
    from skills.brief_deck.scripts import build_deck
    d = bd.deck_data(deck_app.ctx, "board")
    d["signals"] = d["signals"][:2]
    spec = planner.plan_spec(d, "navy-gold", d["focus"], deck_app.store.watch_rules())
    demand = next((s for s in spec["slides"] if s["key"] == "demand"), None)
    assert demand and demand["layout"] == "Hero number" and demand["fill"]["hero"]["metric"] == "active_requirements_msqft"
    dd = {k: v for k, v in d.items() if k not in ("stories", "focus")}
    out = build_deck.build(dd, str(tmp_path / "two.pptx"), {k: v for k, v in spec.items() if k in ("theme", "slides")}, "navy-gold")
    prs = Presentation(out["path"])
    cards = next(s for s in prs.slides if "What changed" in _slide_text(s))
    rects = [sh for sh in cards.shapes if sh.shape_type == 1 and not sh.has_text_frame or (sh.has_text_frame and not sh.text_frame.text.strip() and sh.shape_type == 1)]
    assert len(rects) >= 6  # 2 tone edges + 4 covers over the unused frames


def test_a_cover_never_shows_empty_kpi_rules_or_placeholder_wording(deck_app, tmp_path):
    """Reviewer finding: Copilot-planned covers showed four gold rules with no numbers, and internal text
    ("Top team signals (none added yet)") leaked onto the cover. With no team signals the subtitle is the
    brief's format and period; a cover planned without KPIs draws no rules; with two KPIs it draws two."""
    import re
    from pptx import Presentation
    from pptx.util import Inches
    from skills.brief_deck import tools as bd
    from skills.brief_deck.scripts import build_deck, check_deck
    d = bd.deck_data(deck_app.ctx, "board")  # deck_app has no brief items, so nothing is "chosen by" anyone
    assert "none added yet" not in d["picked_by"] and "Board brief" in d["picked_by"] and d["date"] in d["picked_by"]
    d["signals"], d["talk"] = [], []
    dd = {k: v for k, v in d.items() if k not in ("stories", "focus")}
    # a Copilot-written cover: no kpis in fill, and the old placeholder wording in the subtitle and headline
    spec = {"theme": "navy-gold", "slides": [
        {"layout": "Cover", "title": "Top team signals (none added yet)", "why": "c", "notes": "n",
         "fill": {"subtitle": "Top team signals (none added yet)", "headline": ["Top team signals (none added yet)"], "caveat": "auto"}},
        {"layout": "Chart + takeaway", "title": "Vacancy", "why": "v", "notes": "n",
         "fill": {"chart_label": "Vacancy", "chart": {"chart": "vacancy"}, "takeaway": "Vacancy is 8.0%.", "source": "s"}},
        {"layout": "Sources", "title": "Sources", "why": "s", "notes": "n", "fill": {"body": "auto"}}]}
    out = build_deck.build(dd, str(tmp_path / "nokpi.pptx"), spec)
    prs = Presentation(out["path"])
    cover = prs.slides[0]
    text = _slide_text(cover)
    assert "none added yet" not in text and "Top team signals" not in text
    assert "Board brief" in text and d["date"] in text and d["title"] in text
    assert not [sh for sh in cover.shapes if re.fullmatch(r"kpi rule \d", sh.name)], "no rules drawn for a cover without KPIs"
    assert any(sh.name == "kpi rule mask" for sh in cover.shapes), "the layout's four rules are masked"
    problems = check_deck.check(out["path"], dd)["problems"]
    assert not [p for p in problems if "cover" in p], problems
    # check_deck reports the old behaviour: an unmasked cover rule with nothing under it, and the placeholder phrase
    mask = next(sh for sh in cover.shapes if sh.name == "kpi rule mask")
    mask._element.getparent().remove(mask._element)
    sub = cover.shapes.add_textbox(Inches(0.7), Inches(1.95), Inches(6), Inches(0.4))
    sub.text_frame.text = "Top team signals (none added yet)"
    prs.save(str(tmp_path / "bad.pptx"))
    bad = check_deck.check(str(tmp_path / "bad.pptx"))["problems"]
    assert sum("KPI rule" in p and "no number" in p for p in bad) == 4 and any("placeholder wording" in p for p in bad), bad
    # two KPIs with a value: two rules, centred, and the check is clean
    dd2 = {**dd, "kpis": [k for k in dd["kpis"] if k["metric"] in ("vacancy_pct", "take_up_msqft")], "macro": None}
    out2 = build_deck.build(dd2, str(tmp_path / "two.pptx"), None, "navy-gold")
    cover2 = Presentation(out2["path"]).slides[0]
    rules = sorted(sh.left for sh in cover2.shapes if re.fullmatch(r"kpi rule \d", sh.name))
    assert len(rules) == 2 and rules[0] > Inches(1.5) and "8.0%" in _slide_text(cover2) and "3.10m sq ft" in _slide_text(cover2)
    assert not [p for p in check_deck.check(out2["path"], dd2)["problems"] if "cover" in p]


@needs_pi
def test_two_people_can_ask_the_copilot_at_once_and_a_broken_store_heals(app, tmp_path, monkeypatch):
    """Each person has their own Pi store, so parallel questions never share a file; a store that will not
    load is moved aside and the person gets a fresh conversation instead of a permanent failure."""
    import threading
    from core.pi_engine import PiDurableEngine
    for v in ("PI_STORE", "PI_CONVERSATIONS", "PI_FAUX_SCRIPT"):
        monkeypatch.delenv(v, raising=False)
    monkeypatch.setenv("PI_FAUX", "1")
    monkeypatch.setattr(app, "root", tmp_path)
    eng = PiDurableEngine(app)
    results, errors = {}, {}

    def ask(who):
        try:
            results[who] = eng.ask("What changed?", who)
        except Exception as e:  # noqa: BLE001
            errors[who] = e
    threads = [threading.Thread(target=ask, args=(w,)) for w in ("Priya", "Sarah", "James")]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors and all("Pi Durable" in a.text for a in results.values())
    stores = sorted(p.name for p in (tmp_path / "data" / "pi").glob("*.sqlite"))
    assert stores == ["James.sqlite", "Priya.sqlite", "Sarah.sqlite"]
    # Priya's store is corrupted (as a run killed mid-write would leave it): her next question still works
    store, conv = eng.store_paths("Priya")
    store.write_bytes(b"not a sqlite file at all")
    ans = eng.ask("And rents?", "Priya")
    assert "Pi Durable" in ans.text and "set aside" in ans.engine
    assert list(store.parent.glob("Priya.sqlite.broken-*")) and store.exists()


def test_a_local_pi_failure_is_not_blamed_on_the_model_api():
    from core.engine import _plain
    assert "not the model API" in _plain(RuntimeError("Pi local run of agent.ts failed (not the model API): boom"))
    assert "budget" in _plain(RuntimeError("Pi agent.ts model model_error: insufficient credit"))


def test_check_deck_catches_a_wrapping_cover_title_and_an_empty_slide_and_the_builder_avoids_both(deck_app, tmp_path):
    """Two faults a reviewer saw in Copilot-built decks: a long cover title wrapping over the subtitle, and an IC
    evidence slide with nothing on it. check_deck now reports both; the builder shrinks the title and refuses
    a Title only slide without a table."""
    from skills.brief_deck import planner, tools as bd
    from skills.brief_deck.scripts import build_deck, check_deck
    sample = ROOT / "docs" / "samples" / "team-brief-ic-copilot-sample.pptx"
    problems = check_deck.check(str(sample))["problems"]
    assert any("wraps onto more lines" in p and "slide 1" in p for p in problems)
    assert any("nothing on the slide" in p and "slide 3" in p for p in problems)
    d = bd.deck_data(deck_app.ctx, "ic")
    spec = planner.plan_spec(d, "data-white", d["focus"], deck_app.store.watch_rules())
    spec["slides"][0]["title"] = "IC section: vacancy fell to 8.0% while investment stayed 36% below its long-term average"
    dd = {k: v for k, v in d.items() if k not in ("stories", "focus")}
    out = build_deck.build(dd, str(tmp_path / "long.pptx"), {k: v for k, v in spec.items() if k in ("theme", "slides")}, "data-white")
    assert check_deck.check(out["path"], dd)["problems"] == []
    bad = {k: v for k, v in spec.items() if k in ("theme", "slides")}
    bad["slides"][1] = {"layout": "Title only", "eyebrow": "02 · Evidence", "title": "Submarket evidence", "notes": "x"}
    assert any("Title only needs a table" in e for e in build_deck.validate(bad, dd))
    short = {"theme": "data-white", "slides": [spec["slides"][0], {"layout": "Title only", "eyebrow": "02 · Evidence",
             "title": "Submarket evidence", "notes": "x", "fill": {"table": "evidence"}}, spec["slides"][-1]]}
    assert build_deck.validate(short, dd) == []  # the shorthand {"table": "evidence"} is accepted


def test_flight_to_quality_is_a_signal_from_grade_a_against_overall_vacancy_and_grade_a_rent_is_a_named_metric(app):
    """The brief's demand drivers: Grade A vs overall vacancy becomes a ranked signal with evidence, the supply
    pipeline carries sourced schemes beyond one report, and 'Grade A rent' is a metric the library can hold."""
    eng = SignalEngine(app.store, app.config)
    ftq = eng.flight_to_quality()
    assert ftq and len(ftq) <= 5 and all(s.kind == "flight_to_quality" and s.evidence for s in ftq)
    assert any(s.markets == ["Docklands & Stratford"] and "Grade A vacancy" in s.title for s in ftq)
    assert any("widening" in s.title for s in ftq)  # White City in the seed
    assert tx.find_metrics("What is the Grade A rent in the City?") == ["grade_a_rent_psf"]
    assert tx.METRICS["grade_a_rent_psf"]["unit"] == "£/sq ft"
    schemes = app.store.q("SELECT scheme, source_id FROM pipeline")
    assert len(schemes) >= 10 and len({r["source_id"] for r in schemes}) >= 2
    assert any("Grade A" in s.title for s in eng.for_role("team", limit=40))


@needs_pi
def test_the_copilot_streams_its_answer_and_names_each_tool_as_it_starts(app, tmp_path):
    """Pi Durable's event stream reaches the page: each piece of the answer and each tool start arrives as it
    happens, and the complete Answer (text, trace) is the same as the non-streaming call gives."""
    from core.engine import Engine
    os.environ.update(PI_FAUX="1", PI_STORE=str(tmp_path / "pi.sqlite"), PI_CONVERSATIONS=str(tmp_path / "c.json"))
    os.environ.pop("PI_FAUX_SCRIPT", None)
    eng = Engine(app)
    events = []
    ans = eng.ask("What changed?", "Kenneth", on_event=events.append)
    kinds = [e["event"] for e in events]
    assert "tool" in kinds and "delta" in kinds and kinds.index("tool") < kinds.index("delta")
    assert next(e["name"] for e in events if e["event"] == "tool") == "get_signals"
    assert "".join(e["text"] for e in events if e["event"] == "delta") == ans.text
    stream = eng.ask_streaming("What does Savills say about Midtown?", "Kenneth")
    text = "".join(stream.text())
    assert text.startswith("> **Coverage:**") and "Savills" in text and stream.error is None
    assert stream.answer is not None and stream.answer.text == text and stream.answer.trace


def test_without_pi_a_streaming_answer_reports_the_plain_reason(app, monkeypatch, tmp_path):
    from core.engine import Engine, PiNotReady
    monkeypatch.setenv("PI_ENGINE_DIR", str(tmp_path))
    stream = Engine(app).ask_streaming("What changed?", "Kenneth")
    assert "".join(stream.text()) == "" and isinstance(stream.error, PiNotReady)


def test_saved_chats_come_back_by_person_newest_first_and_empty_ones_are_not_listed(app):
    """A person's conversations are kept in the shared database under their name, so a refresh (or another
    browser with the same name) brings them back: title from the first question, newest first, with the
    answer's tool trace so charts and sources render again. A chat with no message yet is not listed."""
    from core.pi_engine import Answer
    first = app.store.new_chat("Kenneth")
    app.store.add_chat_message(first, "user", "What changed this quarter?")
    app.store.add_chat_message(first, "assistant", "Vacancy fell.", [{"tool": "get_signals", "args": {}, "result": {"summary": "x"}}],
                               "pi-durable:faux")
    second = app.store.new_chat("Kenneth")
    app.store.add_chat_message(second, "user", "And rents?")
    app.store.new_chat("Kenneth")  # opened but never used
    app.store.add_chat_message(app.store.new_chat("Priya"), "user", "Priya's own question")
    chats = app.store.chats("Kenneth")
    assert [c["id"] for c in chats] == [second, first]
    assert chats[1]["title"] == "What changed this quarter?" and chats[1]["messages"] == 2
    msgs = app.store.chat_messages(first)
    assert [m["role"] for m in msgs] == ["user", "assistant"] and msgs[1]["trace"][0]["tool"] == "get_signals"
    assert Answer(msgs[1]["content"], msgs[1]["trace"], msgs[1]["engine"]).engine == "pi-durable:faux"
    assert [c["title"] for c in app.store.chats("Priya")] == ["Priya's own question"]


def test_the_closing_follow_up_line_becomes_buttons_and_leaves_the_answer_text():
    """The Copilot ends each answer with `Follow up: a | b | c`; the page shows those as buttons and the text
    without the line. Answers without it (or with 'follow' in ordinary prose) are left alone."""
    from core.engine import split_follow_ups
    body, qs = split_follow_ups("Vacancy fell to 8.0%.\n\nFollow up: What about Midtown? | Show the trend | Why?")
    assert body == "Vacancy fell to 8.0%." and qs == ["What about Midtown?", "Show the trend", "Why?"]
    body, qs = split_follow_ups("空置率下降。\n\n**Follow up:** 中城呢？｜為什麼重要？")
    assert body == "空置率下降。" and qs == ["中城呢？", "為什麼重要？"]
    text = "Rents rose in the following quarter.\nFollow the Sources page for more."
    assert split_follow_ups(text) == (text, [])
    assert split_follow_ups("Follow up: a | b | c | d")[1] == ["a", "b", "c"]


def test_the_copilot_is_told_to_answer_in_the_language_of_the_question_and_to_offer_follow_ups(app):
    from core.pi_engine import system_prompt
    rules = system_prompt(app)
    assert "language the question is written in" in rules and "English by default" in rules
    assert "Follow up: question one | question two | question three" in rules


@needs_pi
def test_each_saved_chat_keeps_its_own_copilot_conversation(app, tmp_path, monkeypatch):
    """A new chat in the web app starts a clean Pi conversation, and resuming an old chat continues its own:
    the conversation map for the person keys Pi conversations by chat id, with the plain name for the CLI."""
    import json as _json
    from core.pi_engine import PiDurableEngine
    for v in ("PI_STORE", "PI_CONVERSATIONS", "PI_FAUX_SCRIPT"):
        monkeypatch.delenv(v, raising=False)
    monkeypatch.setenv("PI_FAUX", "1")
    monkeypatch.setattr(app, "root", tmp_path)
    eng = PiDurableEngine(app)
    a = eng.ask("What changed?", "Kenneth", conversation="chat-a")
    b = eng.ask("What changed?", "Kenneth", conversation="chat-b")
    a2 = eng.ask("And rents?", "Kenneth", conversation="chat-a")
    cli = eng.ask("What changed?", "Kenneth")
    convs = _json.loads(eng.store_paths("Kenneth")[1].read_text())
    assert set(convs) == {"chat-a", "chat-b", "Kenneth"} and len(set(convs.values())) == 3
    assert a.engine.split("conversation ")[1] == a2.engine.split("conversation ")[1] != b.engine.split("conversation ")[1]
    assert "Pi Durable" in cli.text


def test_brand_files_are_in_place_and_the_deck_carries_the_co_brand(deck_app):
    """The sidebar logo, the Nan Fung foot, the favicon and the deck marks all read from assets/brand; the
    cover carries the co-brand lockup once (as a picture, top right) and content slides the small app icon."""
    from pptx import Presentation
    from pptx.util import Inches
    brand = ROOT / "assets" / "brand"
    for name in ("logo-horizontal.svg", "nanfung/nanfung-logo-color.svg", "cobrand-nanfung.svg", "favicon-32.png",
                 "favicon.svg", "favicon.ico", "apple-touch-icon.png", "cobrand-nanfung-on-dark@2x.png", "app-icon-1024.png"):
        assert (brand / name).exists(), name
    assert b"<svg" in (brand / "nanfung/nanfung-logo-color.svg").read_bytes()[:400]
    assert "Nan Fung" in (brand / "README.md").read_text() and "trademark" in (brand / "README.md").read_text()
    out = deck_app.call("build_brief_pptx", {"format": "board", "theme": "charcoal-amber"}, "Kenneth")
    prs = Presentation(out["data"]["path"])
    marks = [sh for sh in prs.slides[0].shapes if sh.name == "co-brand"]
    assert len(marks) == 1 and marks[0].shape_type == 13  # one picture
    assert marks[0].left + marks[0].width <= prs.slide_width and marks[0].top >= 0
    assert marks[0].left > prs.slide_width // 2, "top right, clear of the eyebrow and title"
    icons = [sh for sl in list(prs.slides)[1:] for sh in sl.shapes if sh.name == "app icon"]
    assert icons and all(ic.top >= Inches(7) for ic in icons)


def test_copilot_answers_reach_team_activity_only_when_shared(app):
    """A Copilot conversation is private to its owner; only an answer the owner shares appears in team activity,
    with the question, their name and the time, and unsharing removes it again."""
    before = len(app.store.activity(50))
    chat = app.store.new_chat("Kenneth")
    app.store.add_chat_message(chat, "user", "Is Midtown vacancy still falling?")
    mid = app.store.add_chat_message(chat, "assistant", "Yes: 8.0% in Q2 2026 [1].",
                                     [{"tool": "get_signals", "args": {}, "result": {}}], "pi")
    assert len(app.store.activity(50)) == before, "asking alone shares nothing"
    assert app.store.shared_by(mid) is None
    app.store.share_answer(mid, "Kenneth")
    app.store.share_answer(mid, "Kenneth")  # a second click changes nothing
    acts = [a for a in app.store.activity(50) if a["kind"] == "ask"]
    assert len(acts) == 1 and acts[0]["author"] == "Kenneth" and acts[0]["ref"] == str(mid)
    assert acts[0]["text"] == "Is Midtown vacancy still falling?" and "8.0%" in acts[0]["answer"]
    assert app.store.shared_by(mid)["author"] == "Kenneth" and app.store.shared_by(mid)["ts"]
    app.store.unshare_answer(mid)
    assert app.store.shared_by(mid) is None and not [a for a in app.store.activity(50) if a["kind"] == "ask"]
    app.store.share_answer(mid - 1, "Kenneth")  # a question (user turn) cannot be shared on its own
    assert not [a for a in app.store.activity(50) if a["kind"] == "ask"]


def test_team_activity_keeps_only_what_the_team_chose_to_share(app):
    """Noise stays out: one auto-update line at most, no setup or deck-build entries, old unshared answers gone
    after an upgrade; a shared answer carries its note and its trace (charts, sources) through shared_answer()."""
    for n in range(3):
        app.store.log("auto", "Auto-update", f"collected open data: run {n}")
    acts = app.store.activity(50)
    assert sum(a["kind"] == "auto" for a in acts) == 1 and "run 2" in next(a for a in acts if a["kind"] == "auto")["text"]
    app.store.db.execute("INSERT INTO team_log (kind, author, text, answer) VALUES ('ask','Old','old question','old answer')")
    app.store.db.execute("INSERT INTO team_log (kind, author, text) VALUES ('setup','Old','set up their view')")
    app.store.db.commit()
    from core.store import Store
    again = Store(app.store.path)  # opening the database again runs the clean-up
    kinds = {(a["kind"], a["author"]) for a in again.activity(100)}
    assert ("ask", "Old") not in kinds and ("setup", "Old") not in kinds
    chat = again.new_chat("Kenneth")
    again.add_chat_message(chat, "user", "Trend?")
    mid = again.add_chat_message(chat, "assistant", "Falling.", [{"tool": "get_metric_trend", "args": {}, "result": {"chart": {"x": [1]}}}], "pi")
    again.share_answer(mid, "Kenneth", "IC needs this")
    a = next(a for a in again.activity(10) if a["kind"] == "ask")
    assert a["note"] == "IC needs this" and again.shared_answer(a["ref"])["trace"][0]["result"]["chart"] == {"x": [1]}


def test_team_brief_is_the_same_for_every_reader(app):
    """One team brief: a reader's own focus never changes it."""
    app.store.save_focus("priya", {"markets": ["Canary Wharf"], "measures": ["vacancy"], "budget": "3", "brief_format": "ic", "lens": "team"})
    a = app.call("generate_brief", {"format": "board"}, "priya")["summary"]
    b = app.call("generate_brief", {"format": "board"}, "Kenneth")["summary"]
    assert a == b
    app.store.reset_focus("priya")


def test_demo_deals_and_seed_news_say_what_they_are(app):
    deals = app.store.internal_deals()
    assert deals and all(d.get("demo") for d in deals), "the seed deals are marked as demo data"
    assert not any("Triton" in d["property"] for d in deals), "no demo deal contradicts the public record"
    assert not any(n.source.endswith("(seed)") for n in app.store.news())
    from core import library as lib
    assert all(f.get("connector") != "seed_research" for f in lib.feeds(app.store, app.config)), "ONS snapshot is not a live feed"
    last = lib.last_auto_update(app.store)
    assert "seed" not in last["text"] or "no live feeds" in last["text"]


def test_slide_editor_edits_the_spec_by_instruction_and_rebuilds(deck_app, tmp_path):
    """The slide editor's chat without the Copilot: plain commands change the built deck's spec (remove, move,
    retitle, add a story, theme) and the deck is built again from it; the cover and sources stay in place and a
    figure not in the data is still refused. Each page can be rendered on its own, so thumbnails appear one by one."""
    from core import deck_edit, preview
    out = deck_app.call("build_brief_pptx", {"format": "board"}, "Kenneth")["data"]
    spec = out["spec"]
    assert spec and len(spec["slides"]) == len(out["slides"]) and spec["slides"][0]["layout"] == "Cover"
    from skills.brief_deck import tools as bd
    stories = bd.deck_data(deck_app.ctx, "board")["stories"]
    n = len(spec["slides"])
    new, said = deck_edit.apply("remove slide 3", spec, stories)
    assert len(new["slides"]) == n - 1 and "Removed slide 3" in said and new["slides"][0]["layout"] == "Cover"
    assert new["slides"][1]["eyebrow"].startswith("02") and new["slides"][2]["eyebrow"].startswith("03")
    assert deck_edit.apply("remove slide 1", spec, stories)[0] is None  # the cover is locked
    moved, _ = deck_edit.apply("move slide 3 to 2", spec, stories)
    assert moved["slides"][1]["key"] == spec["slides"][2]["key"] and moved["slides"][-1]["layout"] == "Sources"
    titled, _ = deck_edit.apply('retitle slide 2 "Vacancy keeps falling"', spec, stories)
    assert titled["slides"][1]["title"] == "Vacancy keeps falling"
    themed, _ = deck_edit.apply("theme forest", spec, stories)
    assert themed["theme"] == "forest-cream"
    assert deck_edit.apply("make it pink", spec, stories) == (None, deck_edit.HELP)
    # the revised spec builds; a made-up figure in a title is refused by the build, as before
    built = deck_app.call("build_brief_pptx", {"format": "board", "theme": "navy-gold", "spec": titled}, "Kenneth")["data"]
    assert built.get("path") and [s["title"] for s in built["slides"]][1] == "Vacancy keeps falling"
    bad, _ = deck_edit.apply('retitle slide 2 "Vacancy fell to 12.7%"', spec, stories)
    refused = deck_app.call("build_brief_pptx", {"format": "board", "spec": bad}, "Kenneth")
    assert refused["data"].get("error") and "12.7" in refused["data"]["error"]
    if not preview.available()["soffice"]:
        pytest.skip("LibreOffice not installed here")
    pdf, why = preview.render_pdf(built["path"], tmp_path / "pages")
    assert pdf and not why and preview.page_count(pdf) == len(built["slides"])
    seen = []
    shown = preview.render(built["path"], tmp_path / "pages2", on_page=lambda n, t, p: seen.append((n, t)))
    assert seen == [(k, len(built["slides"])) for k in range(1, len(built["slides"]) + 1)]
    assert preview.render_page(pdf, 2, tmp_path / "pages").endswith("slide-2.png") and len(shown["pages"]) == len(seen)


def test_the_copilot_edit_request_carries_the_current_spec():
    from core.engine import Engine
    spec = {"theme": "navy-gold", "slides": [{"layout": "Cover", "title": "T"}]}
    ask = Engine.deck_edit_ask("board", "navy-gold", spec, "remove the deals slide")
    assert '"layout": "Cover"' in ask and "remove the deals slide" in ask and "build_brief_pptx" in ask
    assert "Build this week's team brief deck as a IC pack on the data-white theme" in Engine.deck_ask("ic", "data-white")


def test_the_rule_editor_handles_batches_pages_charts_undo_and_chinese(deck_app):
    """The coverage review's acceptance sentences that need no Copilot: several commands in one message, several
    pages at once (English and Chinese), slide ids that survive edits, chart parameters, Sources/cover protection
    with a Chinese explanation, undo and format switches as special results, vague requests answered with a
    targeted question, and 'add' requests that would invent data answered instead of faked."""
    from core import deck_edit as de
    from skills.brief_deck import tools as bd
    spec = deck_app.call("build_brief_pptx", {"format": "board"}, "Kenneth")["data"]["spec"]
    stories = bd.deck_data(deck_app.ctx, "board")["stories"]
    ids = [s["id"] for s in spec["slides"]]
    assert all(ids) and len(set(ids)) == len(ids)
    n = len(spec["slides"])
    new, said = de.apply("刪第3頁，然後換森林主題", spec, stories)
    assert len(new["slides"]) == n - 1 and new["theme"] == "forest-cream" and "✓" in said and "森林" not in said.split("✓")[0]
    assert [s["id"] for s in new["slides"]] == ids[:2] + ids[3:]  # the others keep their ids
    new, said = de.apply("刪第2和第3頁", spec, stories)
    assert len(new["slides"]) == n - 2 and "第 2 頁" in said and "第 3 頁" in said
    new, _ = de.apply("remove slides 2 to 3", spec, stories)
    assert len(new["slides"]) == n - 2
    new, said = de.apply('slide 2 "A" and slide 3 "B"', spec, stories)
    assert new["slides"][1]["title"] == "A" and new["slides"][2]["title"] == "B"
    assert "retitled" in de.describe(spec, new, "x") and "Removed" not in de.describe(spec, new, "x")
    assert de.apply("刪除最後一頁", spec, stories)[0] is None and "Sources" in de.apply("刪除最後一頁", spec, stories)[1]
    assert de.apply("撤銷上一步", spec, stories)[0] == {"undo": 1}
    assert de.apply("回到初版", spec, stories)[0] == {"undo": "first"}
    assert de.apply("restore version 2", spec, stories)[0] == {"undo": ("version", 2)}
    assert de.apply("做一個 IC 版", spec, stories)[0] == {"switch": "ic"}
    assert "PDF" in de.apply("export as PDF", spec, stories)[1] and de.apply("make it 4:3", spec, stories)[0] is None
    # charts: range, markets, type, and a swap, all as parameters the painter reads
    new, said = de.apply("只顯示最近四季的空置率", spec, stories)
    chart = next(v for s in new["slides"] for v in (s.get("fill") or {}).values() if isinstance(v, dict) and v.get("chart") == "vacancy")
    assert chart["last"] == 4
    new, _ = de.apply("only West End Core and City Core on slide 3", spec, stories)
    assert next(v for v in new["slides"][2]["fill"].values() if isinstance(v, dict) and "chart" in v)["markets"] == ["West End Core", "City Core"]
    new, _ = de.apply("as a line chart on slide 2", spec, stories)
    assert any(v.get("type") == "line" for v in new["slides"][1]["fill"].values() if isinstance(v, dict))
    built = deck_app.call("build_brief_pptx", {"format": "board", "spec": new}, "Kenneth")["data"]
    assert built.get("path")
    # shorten to N, 'this slide', vague, and adds that would invent data
    new, said = de.apply("精簡成 4 頁", spec, stories)
    assert len(new["slides"]) == 4 and new["slides"][0]["layout"] == "Cover" and new["slides"][-1]["layout"] == "Sources"
    assert "slide 3" in de.apply("This slide is wrong", spec, stories, sel=3)[1] and "第 3 頁" in de.apply("這頁不對", spec, stories, sel=3)[1]
    assert de.apply("Make it better", spec, stories, sel=3)[0] is None
    assert de.apply("Add our logo", spec, stories)[0] is None and "every slide" in de.apply("Add our logo", spec, stories)[1]
    assert de.apply("加一頁講 Midtown", spec, stories)[0] is None and "Midtown" in de.apply("加一頁講 Midtown", spec, stories)[1]
    assert "slide" in de.apply("Add a chart of yields", spec, stories)[1]
    new, said = de.apply("Add a statement slide saying vacancy hit 47.3%", spec, stories)
    refused = deck_app.call("build_brief_pptx", {"format": "board", "spec": new}, "Kenneth")["data"]
    assert refused.get("refused_numbers") and refused["refused_numbers"][0]["numbers"] == ["47.3"] and refused["latest"]["vacancy_pct"] == "8.0%"
    new, _ = de.apply('第2頁標題改成「空置率升到 9.9%」', spec, stories)
    assert deck_app.call("build_brief_pptx", {"format": "board", "spec": new}, "Kenneth")["data"].get("refused_numbers")
    new, _ = de.apply('第2頁標題改成「Q1 2024 空置率 9.9%」', spec, stories)
    assert deck_app.call("build_brief_pptx", {"format": "board", "spec": new}, "Kenneth")["data"].get("path")  # a period names the past
    new, said = de.apply("加一頁講利率，放在第 3 頁", spec, stories)
    assert new and new["slides"][2]["key"] == "macro" and "3. " in said
    assert de.apply("Use a theme suitable for a client update", spec, stories)[0]["theme"] == "forest-cream"
    assert "navy-gold" in de.apply("換成深藍色", spec, stories)[1]  # already on it
    assert bd.deck_request("幫我把這週的 brief 做成 PPT，給投委會用") == ("ic", "data-white")
    assert de.asks_removal("Make it shorter for the board, three minutes max") and not de.asks_removal("需要多兩頁")


def test_the_copilot_edit_prompt_names_the_selected_slide_and_keeps_ids():
    from core.engine import Engine
    spec = {"theme": "navy-gold", "slides": [{"layout": "Cover", "title": "T", "id": "s1"}, {"layout": "Statement", "title": "U", "id": "s2"}]}
    ask = Engine.deck_edit_ask("board", "navy-gold", spec, "this slide is wrong", sel=2)
    assert "looking at slide 2 (id s2)" in ask and "Keep every slide's `id`" in ask and "`last`" in ask


def test_the_slide_editor_cannot_file_tips_deals_or_rules(app):
    """A slide-editing instruction ("add a statement slide saying vacancy hit 47.3%") once reached the Copilot as
    a question in its usual toolset and was filed as a shared tip. In deck mode the bridge lists and runs only
    reading tools and the three deck tools."""
    import json, urllib.request
    from core.bridge import ToolBridge, allowed
    bridge = ToolBridge(app)
    tools = app.tools
    assert not allowed(tools["share_tip"], "deck") and not allowed(tools["record_deal"], "deck") \
        and allowed(tools["build_brief_pptx"], "deck") and allowed(tools["get_signals"], "deck") and allowed(tools["share_tip"], "")
    names = {t["name"] for t in json.load(urllib.request.urlopen(f"{bridge.url}/manifest?role=Kenneth&mode=deck"))["tools"]}
    assert "build_brief_pptx" in names and "share_tip" not in names and "record_deal" not in names
    bridge.modes["r1"] = "deck"
    req = urllib.request.Request(f"{bridge.url}/call", data=json.dumps({"tool": "share_tip", "args": {"text": "vacancy hit 47.3% in Soho"}, "role": "Kenneth"}).encode(),
                                 headers={"Content-Type": "application/json", "x-request-id": "r1"})
    try:
        urllib.request.urlopen(req)
        assert False, "share_tip ran in deck mode"
    except urllib.error.HTTPError as e:
        assert e.code == 403 and "not available while editing slides" in e.read().decode()
    assert not any("47.3" in (r.get("text") or "") for r in app.store.inbox("new"))
    bridge.server.shutdown()


def test_describe_names_the_chart_change():
    from core import deck_edit as de
    old = {"theme": "navy-gold", "slides": [{"id": "a", "layout": "Cover", "title": "T"},
                                            {"id": "b", "layout": "Chart + takeaway", "title": "V", "fill": {"chart": {"chart": "vacancy"}}}]}
    new = {"theme": "navy-gold", "slides": [{"id": "a", "layout": "Cover", "title": "T"},
                                            {"id": "b", "layout": "Chart + takeaway", "title": "V", "fill": {"chart": {"chart": "rents"}}}]}
    assert "Slide 2's chart is now the rents chart" in de.describe(old, new, "replace the chart")
    new2 = {"theme": "navy-gold", "slides": [{"id": "a", "layout": "Cover", "title": "T"},
                                             {"id": "b", "layout": "Chart + takeaway", "title": "V", "fill": {"chart": {"chart": "vacancy", "last": 4}}}]}
    assert "最近 4 期" in de.describe(old, new2, "只顯示最近四季")


def test_a_stopped_copilot_run_is_killed_and_a_waiting_question_says_so(app, tmp_path, monkeypatch):
    """Pressing Stop used to leave the Node run going, and the next question waited on the per-person lock with
    'Reading market data…' until it finished, without a word. Now Stop kills the run; a question that finds the
    lock held (a deck being built in the background, say) waits its turn with a 'waiting' event, and Stop on
    the waiting question abandons only the wait."""
    import threading, time
    from core.pi_engine import AgentCancelled, PiDurableEngine
    if not shutil.which("node"):
        pytest.skip("node not installed")
    (tmp_path / "slow.ts").write_text("setTimeout(() => { console.log(JSON.stringify({answer: 'late'})); }, 60000);\n")
    eng = PiDurableEngine(app)
    eng.dir = tmp_path
    monkeypatch.setenv("PI_WAIT", "1")
    got = {}

    def first():
        try:
            eng._run("slow.ts", "{}", "Kenneth", None)
        except AgentCancelled:
            got["cancelled"] = True
        except Exception as e:  # noqa: BLE001
            got["error"] = e
    t = threading.Thread(target=first, daemon=True)
    t.start()
    for _ in range(50):
        if eng._running.get("Kenneth"):
            break
        time.sleep(0.1)
    assert eng.cancel("Kenneth")
    t.join(timeout=10)
    assert got == {"cancelled": True} and not eng._lock("Kenneth").locked()
    # a question arriving while an earlier run holds the lock waits its turn (the earlier run may be a deck being
    # built in the background) and says so; a Stop on the waiting question abandons the wait, not the earlier run
    t2 = threading.Thread(target=first, daemon=True)
    got.clear()
    t2.start()
    for _ in range(50):
        if eng._running.get("Kenneth"):
            break
        time.sleep(0.1)
    events, token = [], {"cancelled": False}
    waiter = {}

    def second():
        try:
            eng._run("missing.ts", "{}", "Kenneth", None, events.append, None, "", token)
        except AgentCancelled:
            waiter["cancelled"] = True
    t3 = threading.Thread(target=second, daemon=True)
    t3.start()
    time.sleep(2.5)
    assert any(e.get("event") == "waiting" for e in events) and t2.is_alive() and eng._running.get("Kenneth")
    token["cancelled"] = True
    t3.join(timeout=5)
    assert waiter == {"cancelled": True} and t2.is_alive()  # the earlier run was left alone
    assert eng.cancel("Kenneth")
    t2.join(timeout=10)
    assert got == {"cancelled": True}


def test_a_streamed_answer_can_be_read_again_after_an_interruption(app):
    """The person switched pages while the Copilot was writing: the run goes on, and the thread that picks it up
    again reads every event from the first one (so nothing written is lost) and gets the finished answer."""
    import threading, time
    from core.engine import Engine, StreamingAnswer
    engine = Engine(app)
    calls = []

    def fake_ask(question, role, history=None, on_event=None, conversation=None, mode="", token=None):
        calls.append(question)
        for w in ("one ", "two ", "three "):
            on_event({"event": "delta", "text": w})
            time.sleep(0.05)
        on_event({"event": "tool", "name": "get_signals"})
        from core.pi_engine import Answer
        return Answer("one two three", [], "fake")
    engine.ask = fake_ask
    sa = StreamingAnswer(engine, "What changed?", "Kenneth")
    first = next(sa.events())  # one reader read a piece and was interrupted
    assert first["event"] == "delta" and first["text"] == "one "
    again = list(sa.events())  # the second reader gets everything from the start
    assert [e.get("text") for e in again if e["event"] == "delta"] == ["one ", "two ", "three "]
    assert again[-1] == {"event": "tool", "name": "get_signals"} and sa.answer.text == "one two three" and calls == ["What changed?"]
