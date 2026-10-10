"""Second step of the template build: pptxgenjs names every placeholder "Text n", so give each one the
name the layout catalogue (references/layouts.md) uses, drop the sample slides, save <theme>.pptx."""
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

HERE = Path(__file__).parent
LIGHT = ["eyebrow", "title"]
NAMES = {  # placeholders in the order build_template.js defines them
    "Cover": ["eyebrow", "title", "subtitle", "headline", "kpi1", "kpi2", "kpi3", "kpi4", "caveat"],
    "Section": ["eyebrow", "title", "body"],
    "Statement": ["eyebrow", "statement", "attribution"],
    "Hero number": LIGHT + ["hero", "hero_caption", "hero_context", "chart_label", "chart", "source"],
    "Stat row": LIGHT + [f"stat{i}{s}" for i in range(1, 5) for s in ("", "_caption", "_context")]
    + ["chart_label", "chart", "source"],
    "Chart + takeaway": LIGHT + ["chart_label", "chart", "takeaway", "source"],
    "Two charts": LIGHT + ["left_label", "left_chart", "right_label", "right_chart", "source"],
    "Chart + table": LIGHT + ["chart_label", "chart", "table_label", "table", "source"],
    "Numbered pillars": LIGHT + [f"pillar{i}_{s}" for i in range(1, 4) for s in ("head", "body")],
    "Timeline": LIGHT + [f"step{i}_{s}" for i in range(1, 6) for s in ("when", "head", "body")],
    "Signal cards": LIGHT + ["card1", "card2", "card3", "card4", "card5", "card6", "talk"],
    "Sources": LIGHT + ["body"],
    "Title only": LIGHT,
}

for raw in sorted(HERE.glob("*.raw.pptx")):
    prs = Presentation(raw)
    for layout in prs.slide_layouts:
        names = NAMES.get(layout.name)
        if not names:
            continue
        phs = [p for p in layout.placeholders if p.placeholder_format.type != PP_PLACEHOLDER.SLIDE_NUMBER]
        assert len(phs) == len(names), (raw.name, layout.name, len(phs), names)
        for shape, name in zip(phs, names):
            shape.name = name
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)
    out = HERE / raw.name.replace(".raw.pptx", ".pptx")
    prs.save(out)
    raw.unlink()
    print(out.name, [l.name for l in prs.slide_layouts if l.name in NAMES])
