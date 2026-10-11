"""The team brief as a research note: a Word document laid out like a broker's research report (navy cover with
the co-brand lockup, key takeaways, KPI tiles, the key-numbers table, charts, the signals, deals, sources and an
important-information page), and the same document as a PDF converted by LibreOffice, so the two never differ.

The text is the brief exactly as the page shows it (the cached Markdown, so no model is called again); the
charts are drawn from the library by matplotlib. Every figure keeps its [n], which points to the Sources list."""
from __future__ import annotations

import io
import re
import threading
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
BRAND = ROOT / "assets" / "brand"
NAVY, INK, MUTED, GOLD, RULE, PANEL = "1B3A5C", "1F2A44", "6B7385", "B9781A", "D9D6CC", "F3F1EA"
BLUE, GOOD, BAD = "3A62A8", "2E7D4F", "A8322A"
HEAD, BODY, CJK = "Cambria", "Calibri", "Microsoft JhengHei"
REF = re.compile(r"\[(\d+|internal)\]")
_LOCK = threading.Lock()


# ------------------------------------------------------------------------------------------- the brief, parsed
@dataclass
class Brief:
    title: str = ""
    sub: str = ""
    sections: list[tuple[str, list[str]]] = field(default_factory=list)  # (heading, raw lines)
    sources: list[str] = field(default_factory=list)
    caveat: str = ""

    def section(self, name: str) -> list[str]:
        return next((lines for h, lines in self.sections if h.lower().startswith(name.lower())), [])


def parse(md: str) -> Brief:
    b, cur = Brief(), None
    for line in md.splitlines():
        if line.startswith("# "):
            b.title = line[2:].strip()
        elif line.startswith("## "):
            cur = (line[3:].strip(), [])
            b.sections.append(cur)
        elif cur is None:
            if line.strip():
                b.sub = line.strip().strip("_")
        else:
            cur[1].append(line)
    src = b.section("Sources")
    b.sources = [s for s in src if re.match(r"^\[\d+\]", s)]
    b.caveat = " ".join(s.strip("_ ") for s in src if s.startswith("_"))
    b.sections = [(h, lines) for h, lines in b.sections if h != "Sources"]
    return b


def table_rows(lines: list[str]) -> list[list[str]]:
    rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines if ln.strip().startswith("|")]
    return [r for r in rows if not all(set(c) <= set("-: ") for c in r)]


# ----------------------------------------------------------------------------------------------- docx helpers
# WordprocessingML wants the children of these elements in schema order; Word refuses a file that mixes them up
# (LibreOffice does not mind, so the PDF alone would never show it). _tidy puts them back in order before saving.
_ORDER = {
    "tblPr": "tblStyle tblpPr tblOverlap bidiVisual tblStyleRowBandSize tblStyleColBandSize tblW jc tblCellSpacing "
             "tblInd tblBorders shd tblLayout tblCellMar tblLook tblCaption tblDescription",
    "tcPr": "cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark",
    "pPr": "pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs "
           "suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi "
           "adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection "
           "textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange",
    "rPr": "rStyle rFonts b bCs i iCs caps smallCaps strike dstrike outline shadow emboss imprint noProof snapToGrid "
           "vanish webHidden color spacing w kern position sz szCs highlight u effect bdr shd fitText vertAlign rtl cs "
           "em lang eastAsianLayout specVanish oMath",
    "tcBorders": "top start left bottom end right insideH insideV tl2br tr2bl",
    "pBdr": "top left bottom right between bar",
    "tcMar": "top start left bottom end right",
}
_ORDER = {qn(f"w:{k}"): {qn(f"w:{n}"): i for i, n in enumerate(v.split())} for k, v in _ORDER.items()}


def _tidy(root):
    for el in root.iter(*_ORDER):
        rank = _ORDER[el.tag]
        kids = list(el)
        ordered = sorted(kids, key=lambda k: rank.get(k.tag, len(rank)))  # stable: unknowns keep their place at the end
        if ordered != kids:
            for k in kids:
                el.remove(k)
            el.extend(ordered)


def _shade(cell, hex_):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"), sh.set(qn("w:color"), "auto"), sh.set(qn("w:fill"), hex_)
    tcPr.append(sh)


def _margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), str(int(v))), e.set(qn("w:type"), "dxa")
        mar.append(e)
    tcPr.append(mar)


def _borders(table, color=RULE, inside_h=True, outer=False, size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        on = (edge == "insideH" and inside_h) or (outer and edge in ("top", "bottom"))
        e.set(qn("w:val"), "single" if on else "nil")
        if on:
            e.set(qn("w:sz"), str(size)), e.set(qn("w:color"), color)
        b.append(e)
    tblPr.append(b)


def _cell_border(cell, edge, color, size=12):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single"), e.set(qn("w:sz"), str(size)), e.set(qn("w:color"), color)
    b.append(e)


def _fixed(table, widths_mm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    tblPr.append(lay)
    table.autofit = False
    for row in table.rows:
        for cell, w in zip(row.cells, widths_mm):
            cell.width = Mm(w)
    grid = tbl.tblGrid
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths_mm):
        gc.set(qn("w:w"), str(int(Mm(w).twips)))
    # flush with the text margin: no indent and no default side padding (each cell sets its own)
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), "0"), ind.set(qn("w:type"), "dxa")
    tblPr.append(ind)
    mar = OxmlElement("w:tblCellMar")
    for k in ("left", "right"):
        e = OxmlElement(f"w:{k}")
        e.set(qn("w:w"), "0"), e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblPr.append(mar)


def _no_split(table):
    """Rows never break across a page."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        e = OxmlElement("w:cantSplit")
        trPr.append(e)


def _para_border(p, edge="bottom", color=GOLD, size=8, space=4):
    pPr = p._p.get_or_add_pPr()
    pb = pPr.find(qn("w:pBdr"))
    if pb is None:
        pb = OxmlElement("w:pBdr")
    e = OxmlElement(f"w:{edge}")
    e.set(qn("w:val"), "single"), e.set(qn("w:sz"), str(size)), e.set(qn("w:space"), str(space)), e.set(qn("w:color"), color)
    pb.append(e)
    if pb.getparent() is None:
        pPr.append(pb)


def _keep_next(p):
    p.paragraph_format.keep_with_next = True


def _font(run, name=BODY, size=None, color=None, bold=None, italic=None, caps=False, spacing=None):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    rf = rPr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rPr.insert(0, rf)
    for k in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(k), name)
    rf.set(qn("w:eastAsia"), CJK)
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if caps:
        run.font.all_caps = True
    if spacing is not None:
        sp = OxmlElement("w:spacing")
        sp.set(qn("w:val"), str(int(spacing * 20)))
        rPr.append(sp)
    return run


def _raise(run, half_points=5):
    """A reference raised above the line at a size that stays legible (a true superscript shrinks it to ~5pt)."""
    pos = OxmlElement("w:position")
    pos.set(qn("w:val"), str(half_points))
    run._r.get_or_add_rPr().append(pos)
    return run


def _field(p, code, size=8, color=MUTED):
    for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
        r = p.add_run()
        _font(r, size=size, color=color)
        if kind:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)
        elif text == code:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = f" {code} "
            r._r.append(it)
        else:
            r.text = text


def nice_dates(text: str) -> str:
    """2026-10-08 as 8 Oct 2026, the date style of the rest of the note."""
    from datetime import date as _d

    def one(m):
        try:
            return _d.fromisoformat(m.group(0)).strftime("%d %b %Y").lstrip("0")
        except ValueError:
            return m.group(0)
    return re.sub(r"(?<![\w/-])\d{4}-\d{2}-\d{2}(?![\w/-])", one, text)


def rich(p, text, size=10, color=INK, bold=False, italic=False, font=BODY, sup=True):
    """Markdown-ish inline text: **bold**, _italic_, [n] as a raised reference in navy (inline when `sup` is
    False, as in a source line)."""
    text = nice_dates(text)
    for chunk in re.split(r"(\*\*[^*]+\*\*|(?<![\w])_[^_]+_(?![\w])|\[\d+\]|\[internal\])", text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            _font(p.add_run(chunk[2:-2]), font, size, color, True, italic)
        elif chunk.startswith("_") and chunk.endswith("_") and len(chunk) > 2:
            _font(p.add_run(chunk[1:-1]), font, size, MUTED if not bold else color, bold, True)
        elif REF.fullmatch(chunk):
            r = _font(p.add_run(chunk), BODY, max(size * 0.78, 7.5) if sup else size,
                      NAVY if chunk != "[internal]" else GOLD, True)
            if sup:
                _raise(r)
        else:
            _font(p.add_run(chunk), font, size, color, bold, italic)
    return p


def para(doc_or_cell, text="", size=10, color=INK, bold=False, italic=False, font=BODY, after=4, before=0,
         align=None, line=1.15):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), line
    if align:
        p.alignment = align
    if text:
        rich(p, text, size, color, bold, italic, font)
    return p


def eyebrow(doc, n, text):
    p = para(doc, before=14, after=2)
    _font(p.add_run(f"{n:02d}  ·  {text}"), BODY, 8, GOLD, True, caps=True, spacing=1.2)
    _keep_next(p)
    return p


def heading(doc, text, size=16, n=None, label=None, new_page=False):
    if n is not None:
        e = eyebrow(doc, n, label or text)
        if new_page:
            e.paragraph_format.page_break_before = True
            e.paragraph_format.space_before = Pt(0)
    p = para(doc, after=8, before=0 if n is not None else 12, line=1.0)
    _font(p.add_run(text), HEAD, size, NAVY)
    _para_border(p, "bottom", RULE, 4, 4)
    _keep_next(p)
    return p


# ------------------------------------------------------------------------------------------------------ charts
def _short(period: str) -> str:
    m = re.match(r"(Q\d)\s+(\d{4})", period or "")
    return f"{m.group(1)}\n'{m.group(2)[2:]}" if m else period


def _qnum(period: str):
    m = re.match(r"Q(\d)\s+(\d{4})", period or "")
    return int(m.group(2)) * 4 + int(m.group(1)) - 1 if m else None


def on_time_axis(data: list[dict]) -> tuple[list[dict | None], list[str]]:
    """The series with one slot per quarter from the first to the last, None where the library has no figure,
    so a missing quarter shows as a gap instead of two quarters drawn next to each other. Also the missing
    quarters' names, for the caption."""
    nums = [_qnum(d.get("period", "")) for d in data]
    if not data or None in nums:
        return list(data), []
    by = dict(zip(nums, data))
    slots = [by.get(q) for q in range(min(nums), max(nums) + 1)]
    names = [f"Q{q % 4 + 1} {q // 4}" for q in range(min(nums), max(nums) + 1)]
    return [d if d is not None else {"period": n, "value": None} for d, n in zip(slots, names)], \
        [n for d, n in zip(slots, names) if d is None]


def missing_quarters(data: list[dict]) -> list[str]:
    return on_time_axis(data)[1]


def _chart_png(kind, data, unit="", w=6.6, h=3.0) -> bytes:
    """A chart as a PNG in the house palette. The title and source sit outside it, as figure captions."""
    import logging
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib import font_manager, rcParams
    logging.getLogger("matplotlib.font_manager").setLevel(logging.ERROR)
    have = {f.name for f in font_manager.fontManager.ttflist}
    rcParams.update({"font.family": [f for f in ("Calibri", "Carlito", "DejaVu Sans") if f in have] or ["sans-serif"],
                     "font.size": 8.5, "axes.edgecolor": "#" + RULE, "axes.labelcolor": "#" + MUTED,
                     "xtick.color": "#" + MUTED, "ytick.color": "#" + MUTED, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.spines.left": False, "ytick.left": False})
    fig, ax = plt.subplots(figsize=(w, h), dpi=220)
    ax.set_facecolor("white")
    label = {"fontsize": 7.5, "color": "#" + INK, "bbox": {"boxstyle": "square,pad=0.15", "fc": "white", "ec": "none"}}
    if kind in ("line", "column"):
        data, gaps = on_time_axis(data)
        xs = list(range(len(data)))
        ys = [d["value"] for d in data]
        have = [y for y in ys if y is not None]
        lta = [d.get("lta") for d in data if d.get("lta") is not None]
        ax.set_xticks(xs, [_short(d["period"]) for d in data])
        if kind == "line":
            seg = []  # a missing quarter breaks the line rather than being bridged
            for x, y in zip(xs + [None], ys + [None]):
                if y is None:
                    if seg:
                        ax.plot([a for a, _ in seg], [b for _, b in seg], color="#" + NAVY, lw=2.0, marker="o",
                                ms=3.5, zorder=3)
                    seg = []
                else:
                    seg.append((x, y))
            ax.plot(xs[-1:], ys[-1:], color="#" + GOLD, marker="o", ms=6, zorder=4)
            first = next(i for i, y in enumerate(ys) if y is not None)
            for i in (first, len(ys) - 1):
                ax.annotate(f"{ys[i]:.1f}{unit}", (i, ys[i]), textcoords="offset points", xytext=(0, 7), ha="center",
                            fontsize=8, color="#" + INK, fontweight="bold")
            lo, hi = min(have + lta), max(have + lta)
            ax.set_ylim(lo - (hi - lo) * 0.35, hi + (hi - lo) * 0.3)
            for i, y in enumerate(ys):
                if y is None:
                    ax.text(i, lo - (hi - lo) * 0.2, "no\ndata", ha="center", va="center", fontsize=6.5,
                            color="#" + MUTED, style="italic")
        else:
            top = max(have + lta)
            ax.bar([x for x, y in zip(xs, ys) if y is not None], have, width=0.62, zorder=3,
                   color=["#" + BLUE] * (len(have) - 1) + ["#" + GOLD])
            avg = lta[-1] if lta else None
            for i, v in enumerate(ys):
                if v is None:
                    ax.text(i, top * 0.04, "no\ndata", ha="center", va="bottom", fontsize=6.5, color="#" + MUTED,
                            style="italic")
                    continue
                near = avg is not None and abs(v - avg) < top * 0.06  # a label on the dashed line: lift it clear
                ax.text(i, (max(v, avg) + top * 0.035) if near else v + top * 0.015, f"{v:.2f}", ha="center",
                        va="bottom", zorder=5, fontsize=7.5, color="#" + INK,
                        fontweight="bold" if i == len(ys) - 1 else "normal")
            ax.set_ylim(0, top * 1.18)
        if lta:
            ax.axhline(lta[-1], color="#" + (GOLD if kind == "line" else NAVY), lw=1.1, ls=(0, (4, 3)), zorder=2)
            if kind == "line":  # the column chart names its average in the caption, clear of the bars
                ax.text(-0.3, lta[-1], f"Long-term avg {lta[-1]:.1f}{unit}", va="top", ha="left", zorder=5,
                        **{**label, "color": "#" + GOLD})
        ax.margins(x=0.07)
        ax.grid(axis="y", color="#" + RULE, lw=0.5)
        ax.tick_params(axis="x", labelsize=6.8, length=0)
        ax.tick_params(axis="y", labelsize=7)
    else:  # horizontal bars by market
        data = sorted(data, key=lambda d: d["value"])[-12:]
        names = [d["market"] for d in data]
        ys = [d["value"] for d in data]
        moved = [bool(d.get("change")) for d in data]
        ax.barh(names, ys, color=["#" + GOLD if m else "#" + BLUE for m in moved], height=0.62, zorder=3)
        for i, v in enumerate(ys):
            ax.text(v, i, f" {unit}{v:,.2f}" if unit == "£" else f" {v:g}", va="center", fontsize=7.5, color="#" + INK)
        ax.grid(axis="x", color="#" + RULE, lw=0.5)
        ax.set_xlim(0, max(ys) * 1.15)
        ax.tick_params(axis="y", labelsize=7.5, length=0)
        ax.tick_params(axis="x", labelsize=7)
        ax.spines["left"].set_visible(True)
    fig.tight_layout(pad=0.4)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor="white", bbox_inches="tight", pad_inches=0.06)
    plt.close(fig)
    return buf.getvalue()


# ------------------------------------------------------------------------------------------------- the document
def build_docx(md: str, fmt: str, deck: dict, out: Path, team: str = "London office team") -> Path:
    """`md` is the brief as written; `deck` is skills.brief_deck.tools.deck_data (series for the charts)."""
    b = parse(md)
    sources = list(b.sources)

    def cite(title_text: str) -> str:
        """[n] for a source title, added to the Sources list when the text did not cite it."""
        base = re.sub(r"\s*\(Uploaded by.*$", "", re.sub(r"^\[\d+\]\s*", "", title_text)).split(" - http")[0].strip()
        for i, s in enumerate(sources, 1):
            if base and base in s:
                return f"[{i}]"
        sources.append(f"[{len(sources) + 1}] {re.sub(r'^\[\d+\]\s*', '', title_text)}")
        return f"[{len(sources)}]"

    def deck_cite(ref: str) -> str:
        m = re.match(r"\[(\d+)\]", ref or "")
        foot = deck.get("footnotes") or []
        return cite(foot[int(m.group(1)) - 1]) if m and int(m.group(1)) <= len(foot) else ""

    doc = Document()
    styles = doc.styles
    st = styles["Normal"]
    st.font.name, st.font.size = BODY, Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    is_ic = fmt == "ic"
    kind = "IC paper section" if is_ic else "Board summary"
    date_text = (b.sub.split("·")[0].strip() if b.sub else "")
    period = next((r[3] for r in table_rows(b.section("Key numbers"))[1:2]), "")
    report = "London Office Market · " + ("IC Paper" if is_ic else "Weekly Brief")
    doc.core_properties.title = report
    doc.core_properties.author = team
    doc.core_properties.subject = "London office market research note"

    # ---- cover: a full-bleed navy band
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Mm(0)
    sec.header_distance = sec.footer_distance = Mm(0)
    band = doc.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    _fixed(band, [210])
    _borders(band, inside_h=False)
    row = band.rows[0]
    row.height, row.height_rule = Mm(178), WD_ROW_HEIGHT_RULE.EXACTLY
    c = row.cells[0]
    _shade(c, NAVY)
    _margins(c, top=Mm(22).twips, left=Mm(22).twips, right=Mm(22).twips, bottom=Mm(14).twips)
    c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    p = c.paragraphs[0]
    p.add_run().add_picture(str(BRAND / "cobrand-nanfung-on-dark@2x.png"), width=Mm(92))
    para(c, after=0, before=0).add_run().add_break()
    for _ in range(5):
        para(c, after=6)
    p = para(c, after=6)
    _font(p.add_run(f"RESEARCH  ·  {kind.upper()}"), BODY, 9, GOLD, True, spacing=2)
    p = para(c, after=10, line=1.0)
    _font(p.add_run("London Office Market"), HEAD, 34, "FFFFFF")
    p = para(c, after=18, line=1.0)
    _font(p.add_run("Investment committee section" if is_ic else "Weekly brief"), HEAD, 22, "C9D1DE", italic=True)
    p = para(c, after=4)
    _para_border(p, "top", GOLD, 12, 10)
    hl = " ".join(ln for ln in b.section("Headline") if ln.strip())
    first = re.split(r"(?<=\.)\s+", REF.sub("", hl).strip())[0] if hl else ""
    if first:
        rich(p, first, 12.5, "FFFFFF", font=HEAD, italic=True)
    p = para(c, after=0, before=10)
    _font(p.add_run(f"{date_text}   ·   Data to {period}   ·   {team}"), BODY, 9.5, "C9D1DE")

    meta = doc.add_table(rows=1, cols=2)
    _fixed(meta, [126, 84])
    _borders(meta, inside_h=False)
    left, right = meta.rows[0].cells
    for cell in (left, right):
        _margins(cell, top=Mm(16).twips, left=Mm(22).twips, right=Mm(22).twips)
    p = left.paragraphs[0]
    _font(p.add_run("ABOUT THIS NOTE"), BODY, 8, GOLD, True, spacing=1.2)
    para(left, "Built from the signals the team added this week and the figures in the shared library. Every "
               "figure carries a reference [n] to its source, listed at the end.", 9.5, INK, after=6)
    para(left, b.sub.replace("·", " · ") if b.sub else "", 9, MUTED, italic=True, after=0)
    p = right.paragraphs[0]
    _font(p.add_run("CONTENTS"), BODY, 8, GOLD, True, spacing=1.2)
    toc = ["Key takeaways", "Key numbers", "Charts"] + [h for h, _ in b.sections if h not in ("Headline", "Key numbers")] \
        + ["Sources", "Important information"]
    for i, t in enumerate(toc, 1):
        p = para(right, after=1)
        _font(p.add_run(f"{i:02d}  "), BODY, 9, GOLD, True)
        _font(p.add_run(t), BODY, 9, INK)
    p = para(doc, before=0, after=0)
    p.paragraph_format.left_indent = Mm(22)

    # ---- content section with running header and footer
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.top_margin, sec.bottom_margin = Mm(30), Mm(20)
    sec.left_margin = sec.right_margin = Mm(20)
    sec.header_distance, sec.footer_distance = Mm(10), Mm(9)
    sec.header.is_linked_to_previous = False
    sec.footer.is_linked_to_previous = False
    hdr = sec.header
    ht = hdr.add_table(rows=1, cols=2, width=Mm(170))
    _fixed(ht, [80, 90])
    _borders(ht, inside_h=False)
    hl_, hr_ = ht.rows[0].cells
    hl_.paragraphs[0].add_run().add_picture(str(BRAND / "cobrand-nanfung@2x.png"), width=Mm(56))
    hl_.vertical_alignment = hr_.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = hr_.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _font(p.add_run(report.upper()), BODY, 7.5, NAVY, True, spacing=0.8)
    p = para(hr_, after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _font(p.add_run(f"{date_text} · Data to {period}"), BODY, 7.5, MUTED)
    p = hdr.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    _para_border(p, "bottom", GOLD, 8, 1)
    hdr._element.append(p._p)  # the gold rule under the header table
    ft = sec.footer
    p = ft.paragraphs[0]
    _para_border(p, "top", RULE, 4, 4)
    tabs = p.paragraph_format.tab_stops
    tabs.add_tab_stop(Mm(170), alignment=2)
    _font(p.add_run("London Office Market Agent  ·  Prepared for Nan Fung Group  ·  For internal discussion"), BODY, 7.5, MUTED)
    _font(p.add_run("\tPage "), BODY, 7.5, MUTED)
    _field(p, "PAGE", 7.5, NAVY)
    _font(p.add_run(" of "), BODY, 7.5, MUTED)
    _field(p, "NUMPAGES", 7.5, MUTED)

    n = 1
    # ---- key takeaways
    heading(doc, "Key takeaways", 18, n, "Summary")
    box = doc.add_table(rows=1, cols=1)
    _fixed(box, [170])
    _no_split(box)
    _borders(box, inside_h=False)
    cell = box.rows[0].cells[0]
    _shade(cell, PANEL)
    _cell_border(cell, "left", GOLD, 24)
    _margins(cell, top=Mm(4).twips, bottom=Mm(3).twips, left=Mm(6).twips, right=Mm(6).twips)
    sentences = [s for s in re.split(r"(?<=[.!?])\s+(?=[A-Z£\d])", hl.strip()) if s.strip()]
    refs = "".join(REF.findall(hl) and [f"[{r}]" for r in dict.fromkeys(REF.findall(hl))])
    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
    for i, s in enumerate(sentences):
        s = REF.sub("", s).strip()
        p = cell.paragraphs[0] if i == 0 else para(cell)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.left_indent = Mm(5)
        p.paragraph_format.first_line_indent = Mm(-5)
        _font(p.add_run("■  "), BODY, 7, GOLD)
        rich(p, s + (" " + refs if i == len(sentences) - 1 else ""), 10.5, INK, font=BODY)

    # KPI tiles from the first four London rows of the key-numbers table
    rows = table_rows(b.section("Key numbers"))
    head, body = (rows[0], rows[1:]) if rows else ([], [])
    tiles = [r for r in body if r[0] == "London"][:4]
    if tiles:
        para(doc, after=2)
        kt = doc.add_table(rows=1, cols=len(tiles))
        _fixed(kt, [170 / len(tiles) - 0.01] * len(tiles))
        _no_split(kt)
        _borders(kt, inside_h=False)
        for cell, r in zip(kt.rows[0].cells, tiles):
            _cell_border(cell, "top", GOLD, 18)
            _margins(cell, top=Mm(3).twips, left=Mm(2).twips, right=Mm(3).twips)
            p = cell.paragraphs[0]
            _font(p.add_run(r[1].upper()), BODY, 7, MUTED, True, spacing=0.6)
            val = REF.sub("", r[2]).strip()
            p = para(cell, after=0, line=1.0)
            _font(p.add_run(val), HEAD, 17, NAVY)
            for m in REF.findall(r[2]):
                _raise(_font(p.add_run(f"[{m}]"), BODY, 8.5, NAVY, True), 8)
            chg = r[4] if len(r) > 4 else ""
            up = chg.startswith("+") and not chg.startswith("+0.00")
            down = chg.startswith("-")
            p = para(cell, after=0)
            _font(p.add_run(("▲ " if up else "▼ " if down else "● ") + (chg or r[3])), BODY, 8,
                  GOOD if up and "Vacancy" not in r[1] else BAD if (down and "Vacancy" not in r[1]) or (up and "Vacancy" in r[1])
                  else GOOD if down else MUTED)
    # ---- key numbers table
    n += 1
    heading(doc, "Key numbers", 16, n, "Market data")
    if rows:
        widths = [32, 46, 32, 20, 40]
        t = doc.add_table(rows=len(rows), cols=len(head))
        _fixed(t, widths[: len(head)])
        _no_split(t)
        h_ = OxmlElement("w:tblHeader")  # the header row repeats when the table runs onto another page
        t.rows[0]._tr.get_or_add_trPr().append(h_)
        _borders(t, RULE, True, True, 4)
        for ri, r in enumerate(rows):
            for ci, txt in enumerate(r):
                cell = t.rows[ri].cells[ci]
                _margins(cell, top=Mm(1.4).twips, bottom=Mm(1.4).twips, left=Mm(2).twips, right=Mm(2).twips)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                if ci >= 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if ci == 2 else WD_ALIGN_PARAGRAPH.LEFT
                if ri == 0:
                    _shade(cell, NAVY)
                    _font(p.add_run(txt.upper()), BODY, 7.5, "FFFFFF", True, spacing=0.5)
                else:
                    if ri % 2 == 0:
                        _shade(cell, "F7F6F2")
                    color = INK
                    if ci == 4 and txt.startswith(("+", "-")) and not re.match(r"[+-]0\.00", txt):
                        color = GOOD if txt.startswith("+") else BAD
                        if "Vacancy" in r[1]:
                            color = BAD if txt.startswith("+") else GOOD
                    rich(p, txt, 9, color, bold=(ci == 2))
        p = para(doc, before=3, after=0)
        _font(p.add_run("Change is against the previous quarter; pp = percentage points. "
                        "Green and red show whether the move is good or bad for an owner of London offices."), BODY, 7.5, MUTED, italic=True)
    flagged = [ln for ln in b.section("Key numbers") if ln.startswith("_")]
    for f in flagged:
        para(doc, f, 8, BAD, italic=True, after=2)

    # ---- charts
    n += 1
    heading(doc, "Charts", 16, n, "Trends")
    charts = []  # (kind, series, caption, unit, ref of the latest point)
    if deck.get("vacancy"):
        charts.append(("line", deck["vacancy"], "Central London vacancy rate (%)", "%", deck["vacancy"][-1].get("ref")))
    if deck.get("take_up"):
        charts.append(("column", deck["take_up"], "Central London take-up (m sq ft)", "", deck["take_up"][-1].get("ref")))
    if deck.get("rents"):
        charts.append(("barh", deck["rents"], "Prime office rent by submarket (£ per sq ft)", "£",
                       deck["rents"][0].get("ref")))

    def figure(cell_or_doc, i, k, data, caption, unit, ref, width_mm):
        p = para(cell_or_doc, after=3, line=1.0)
        _keep_next(p)
        _font(p.add_run(f"Figure {i}   "), BODY, 8.5, GOLD, True)
        _font(p.add_run(caption), BODY, 8.5, NAVY, True)
        lta = [d.get("lta") for d in data if d.get("lta") is not None]
        note = {"line": "Quarter end; gold dashed line: long-term average",
                "column": "Latest quarter in gold" + (f"; dashed line: long-term average {lta[-1]:.2f}" if lta else ""),
                "barh": "Latest quarter; gold: moved since the previous quarter"}[k]
        if k in ("line", "column") and (gone := missing_quarters(data)):
            note += f". No figure in the library for {' and '.join(gone)}"
        p = para(cell_or_doc, after=4, line=1.0)
        _keep_next(p)
        _font(p.add_run(note), BODY, 7.5, MUTED, italic=True)
        half = width_mm < 120
        png = _chart_png(k, data, unit, w=width_mm / 25.4, h=(2.7 if half else 3.3 if k == "barh" else 2.5))
        p = para(cell_or_doc, after=0, line=1.0)
        _keep_next(p)
        p.add_run().add_picture(io.BytesIO(png), width=Mm(width_mm - (3 if half else 0)))
        p = para(cell_or_doc, after=0, before=2, line=1.0)
        _para_border(p, "top", RULE, 4, 2)
        _font(p.add_run("Source: "), BODY, 7.5, MUTED)
        rich(p, deck_cite(ref) or "the team library", 7.5, NAVY, sup=False)
        if k == "line":
            refs = sorted({deck_cite(d.get("ref")) for d in data if d.get("ref")} - {"", deck_cite(ref)},
                          key=lambda s: int(s[1:-1]))
            if refs:
                _font(p.add_run("; earlier quarters "), BODY, 7.5, MUTED)
                rich(p, "".join(refs), 7.5, NAVY, sup=False)

    pair = [c for c in charts if c[0] in ("line", "column")]
    num = 0
    if len(pair) == 2:  # the two time series side by side, as a research note sets them
        t = doc.add_table(rows=1, cols=3)
        _fixed(t, [82, 6, 82])
        _no_split(t)
        _borders(t, inside_h=False)
        for cell, c in zip((t.rows[0].cells[0], t.rows[0].cells[2]), pair):
            num += 1
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            figure(cell, num, *c, 82)
            cell._tc.remove(cell.paragraphs[0]._p)  # the empty first paragraph python-docx adds
        rest = [c for c in charts if c not in pair]
        para(doc, after=8)
    else:
        rest = charts
    for c in rest:
        num += 1
        figure(doc, num, *c, 170)
        para(doc, after=6)

    # ---- the remaining sections, as written
    for h, lines in b.sections:
        if h in ("Headline", "Key numbers"):
            continue
        n += 1
        heading(doc, h, 16, n, {"What changed": "Signals", "Supports": "Signals", "Risks": "Signals", "Other": "Signals",
                                "Our deals": "Internal", "Market deals": "Evidence", "Market talk": "Unconfirmed",
                                "Macro": "Context"}.get(next((k for k in ("What changed", "Supports", "Risks", "Other",
                                                                          "Our deals", "Market deals", "Market talk",
                                                                          "Macro") if h.startswith(k)), ""), "Detail"))
        for ln in lines:
            s = ln.strip()
            if not s:
                continue
            if s.startswith("- "):
                s = s[2:]
                m = re.match(r"^([▲▼●])\s+\*\*(.+?)\*\*\.?\s*(.*)$", s)
                if m:  # a signal: marker, title, detail [n], _why_
                    arrow, title, rest = m.groups()
                    why = ""
                    if (w := re.search(r"\s_([^_]+)_\s*$", rest)):
                        why, rest = w.group(1), rest[: w.start()]
                    t = doc.add_table(rows=1, cols=2)
                    _fixed(t, [8, 162])
                    _no_split(t)
                    _borders(t, inside_h=False)
                    a, body_c = t.rows[0].cells
                    _margins(a, top=Mm(1).twips)
                    _margins(body_c, top=Mm(0.6).twips, bottom=Mm(2.4).twips, left=Mm(1).twips)
                    _cell_border(body_c, "bottom", RULE, 4)
                    _cell_border(a, "bottom", RULE, 4)
                    pa = a.paragraphs[0]
                    _font(pa.add_run(arrow), BODY, 11, GOOD if arrow == "▲" else BAD if arrow == "▼" else MUTED, True)
                    p = body_c.paragraphs[0]
                    p.paragraph_format.space_after = Pt(2)
                    rich(p, title, 10.5, NAVY, bold=True, font=HEAD)
                    rich(para(body_c, after=2), rest.strip(), 9.5, INK)
                    if why:
                        p = para(body_c, after=0)
                        _font(p.add_run("Why it matters  "), BODY, 7.5, GOLD, True, spacing=0.5)
                        rich(p, why, 9, MUTED, italic=True)
                else:
                    p = para(doc, after=3)
                    p.paragraph_format.left_indent = Mm(5)
                    p.paragraph_format.first_line_indent = Mm(-5)
                    _font(p.add_run("■  "), BODY, 7, GOLD)
                    rich(p, s, 9.5, INK)
            elif s.startswith("_") and s.endswith("_"):
                para(doc, s, 8.5, MUTED, italic=True, after=4)
            else:
                para(doc, s, 10, INK, after=4)

    # ---- sources
    n += 1
    heading(doc, "Sources", 16, n, "References")
    for s in sources:
        m = re.match(r"^\[(\d+)\]\s*(.*)$", s)
        if not m:
            continue
        num, rest = m.groups()
        title_, _, url = rest.partition(" - http")
        p = para(doc, after=3)
        p.paragraph_format.left_indent = Mm(8)
        p.paragraph_format.first_line_indent = Mm(-8)
        _font(p.add_run(f"[{num}]\t"), BODY, 8.5, NAVY, True)
        p.paragraph_format.tab_stops.add_tab_stop(Mm(8))
        _font(p.add_run(nice_dates(title_.strip())), BODY, 8.5, INK)
        if url:
            _font(p.add_run("  http" + url.strip()), BODY, 7.5, BLUE)

    # ---- important information
    n += 1
    heading(doc, "Important information", 16, n, "Disclaimer", new_page=True)
    info = [
        ("About this note", "This research note was generated by the London Office Market Agent, a proof of concept "
         "built for Nan Fung Group's London office team. It summarises the signals the team selected and the figures "
         "in the team's shared library at the date shown."),
        ("Data", (b.caveat or "") + " Market figures are restated from public Knight Frank research reports for "
         "evaluation of this proof of concept only; Knight Frank's terms do not permit the use of their data by AI "
         "systems or for commercial purposes. Bank Rate is from the Bank of England. Internal deals are demonstration "
         "data. Verify every figure against the original source, linked in Sources, before any external use."),
        ("Unconfirmed information", "Market talk shared by colleagues is shown separately and marked unconfirmed; "
         "do not rely on it until a published source confirms it."),
        ("Not advice", "Nothing in this note is investment, legal or tax advice, or an offer or solicitation to buy "
         "or sell any property or security. It is intended for internal discussion within the team and is not for "
         "distribution outside it."),
        ("Trademarks", "The Nan Fung Group name and logo are trademarks of Nan Fung Group, used here only for this "
         "interview proof of concept; their use is not a claim of affiliation or endorsement."),
    ]
    for h, text in info:
        p = para(doc, before=6, after=2)
        _font(p.add_run(h.upper()), BODY, 8, NAVY, True, spacing=0.8)
        para(doc, text, 9, INK, after=4, line=1.25)
    p = para(doc, before=18, after=0)
    p.add_run().add_picture(str(BRAND / "cobrand-nanfung@2x.png"), width=Mm(70))

    _tidy(doc.element)
    for s_ in doc.sections:
        for part in (s_.header, s_.footer):
            if not part.is_linked_to_previous:
                _tidy(part._element)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def build(ctx, md: str, fmt: str, out_dir: Path) -> dict:
    """The .docx and its PDF under out_dir. Returns {"docx", "pdf", "why"}; without LibreOffice there is no PDF
    and `why` says so."""
    from core import preview
    from skills.brief_deck import tools as bd
    deck = bd.deck_data(ctx, fmt)
    name = "london-office-ic-paper" if fmt == "ic" else "london-office-weekly-brief"
    with _LOCK:  # matplotlib's pyplot is not thread-safe, and two pages may ask at once
        docx = build_docx(md, fmt, deck, Path(out_dir) / f"{name}.docx",
                          ctx.config.get("team", {}).get("name", "London office team"))
    pdf, why = preview.render_pdf(docx, out_dir)
    if pdf:
        final = Path(out_dir) / f"{name}.pdf"
        Path(pdf).replace(final)
        pdf = str(final)
    return {"docx": str(docx), "pdf": pdf, "why": why}
